import base64
from io import BytesIO
import unittest
import openpyxl
from unit_ledger_fixture import fixture_bytes,mapping
from project_import import currency_name
from server import parse_file,export_report
from reconcile import reconcile_projects,reconcile_account
from cashflow import calculate_cashflow


def parsed(alias=False):
    d=parse_file({'name':'单位三栏账.xlsx','content':base64.b64encode(fixture_bytes(alias)).decode(),'autoHeader':True})
    d.update(name='单位三栏账.xlsx',mapping=mapping())
    return d


class UnitLedgerTests(unittest.TestCase):
    def test_two_level_headers_split_dates_and_original_columns(self):
        d=parsed()
        self.assertEqual(d['ledgerBooks'],['丙项目','乙项目','甲项目'])
        self.assertEqual(d['ledgerCounts']['details'],7)
        self.assertEqual(d['ledgerCounts']['summary'],3)
        self.assertEqual(d['rowNumbers'][0],9)
        self.assertEqual(d['rows'][0][2],'2026-01-01')
        self.assertEqual(d['rows'][0][7],'100')
        self.assertEqual(d['rows'][0][9],'USD')
        self.assertEqual(d['rows'][2][6],'MVR')

    def test_internal_external_summary_self_and_base_only_are_distinct(self):
        d=parsed();r=reconcile_projects({'datasets':[d],'options':{'days':'unlimited'}})
        self.assertEqual(r['ledgerExclusionCounts'],{'summary':3,'unconfirmed':1,'zeroOriginal':1,'self':1})
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(len(r['differences']),1)
        self.assertFalse(r['invalid'])
        self.assertEqual(r['counts']['rows']+len(r['ledgerExcluded']),10)
        self.assertEqual({r['matched'][0]['left']['row'],r['matched'][0]['right']['row']},{9,10})
        self.assertFalse(any('外部客户' in w and '未提供账簿' in w for w in r['warnings']))

    def test_project_overview_keeps_each_original_currency(self):
        r=reconcile_projects({'datasets':[parsed()],'options':{'days':'unlimited'}})
        a=[v for v in r['unitOverview'] if v['project']=='甲项目']
        self.assertEqual(len(a),2)
        usd=next(v for v in a if v['currency']=='USD');mvr=next(v for v in a if v['currency']=='MVR')
        self.assertEqual(usd['delta'],0)
        self.assertEqual(usd['ownAmount'],10000)
        self.assertEqual(mvr['delta'],1000)
        self.assertEqual(mvr['pendingCount'],1)

    def test_alias_resolves_before_internal_unit_filtering(self):
        r=reconcile_projects({'datasets':[parsed(True)],'aliases':{'旧名乙':'乙项目'},'options':{'days':'unlimited'}})
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(r['ledgerExclusionCounts']['unconfirmed'],1)

    def test_additional_internal_unit_marks_missing_book(self):
        r=reconcile_projects({'datasets':[parsed()],'options':{'days':'unlimited','additionalInternalUnits':['外部客户']}})
        self.assertEqual(r['ledgerExclusionCounts'].get('unconfirmed',0),0)
        self.assertTrue(any(v['project']=='外部客户' and not v['provided'] for v in r['unitOverview']))
        self.assertTrue(any('外部客户' in w and '没有提供账簿' in w for w in r['warnings']))

    def test_unlimited_date_range_ignores_stale_date_inputs(self):
        r=reconcile_projects({'datasets':[parsed()],'options':{'days':'unlimited','dateRangeMode':'unlimited','startDate':'bad','endDate':'1900-01-01'}})
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(r['options']['startDate'],'')
        self.assertEqual(r['options']['endDate'],'')
        self.assertFalse(r['excludedPeriod'])

    def test_workbook_has_discovery_project_overview_and_exclusions(self):
        r=reconcile_projects({'datasets':[parsed()],'options':{'days':'unlimited'}})
        wb=openpyxl.load_workbook(BytesIO(export_report(r)))
        self.assertEqual(wb['内部单位清单'].max_row,4)
        self.assertGreater(wb['按项目核对概览'].max_row,3)
        self.assertEqual(wb['未纳入内部核对'].max_row,7)
        self.assertFalse(any(c.data_type=='f' for s in wb for row in s for c in row));wb.close()

    def test_currency_code_prefix_is_preserved_as_actual_currency(self):
        for raw,wanted in [('USD-美元','USD'),('MVR-马尔代夫拉菲亚','MVR'),('CNY-人民币','CNY'),('LKR-斯里兰卡卢比','LKR'),('RMB','CNY'),('RMB-人民币','CNY')]:
            self.assertEqual(currency_name(raw),wanted)


class UnlimitedDateTests(unittest.TestCase):
    def test_one_to_one_across_years(self):
        from test_extended import ds,row
        a=ds([row(100,other='B',dt='2020-01-01')],project='A');b=ds([row(-100,other='A',dt='2026-09-01')],project='B')
        r=reconcile_projects({'datasets':[a,b],'options':{'days':'unlimited'}})
        self.assertEqual(len(r['matched']),1)
        self.assertIn('日期无限制',r['matched'][0]['reason'])

    def test_one_to_many_without_date_windows(self):
        from test_extended import ds,row
        a=ds([row(100,other='B',dt='2020-01-01')],project='A')
        b=ds([row(-60,other='A',dt='2024-01-01'),row(-40,other='A',dt='2026-09-01')],project='B')
        r=reconcile_projects({'datasets':[a,b],'options':{'days':'unlimited','oneToMany':True}})
        self.assertEqual(r['matched'][0]['matchType'],'one_to_many')
        self.assertEqual(len(r['matched'][0]['rightRows']),2)

    def test_unlimited_does_not_arbitrarily_choose_repeated_amounts(self):
        from test_extended import ds,row
        a=ds([row(100,other='B',dt='2020-01-01')],project='A')
        b=ds([row(-100,other='A',dt='2024-01-01'),row(-100,other='A',dt='2026-09-01')],project='B')
        r=reconcile_projects({'datasets':[a,b],'options':{'days':'unlimited'}})
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']),3)

    def test_bank_account_is_still_required_when_dates_unlimited(self):
        from test_cashflow import ledger,transaction
        a=ledger([transaction('2020-01-01','A','100',account='BANK1')]);b=ledger([transaction('2026-09-01','A','100',account='BANK1')])
        r=reconcile_account({'left':a,'right':b,'options':{'days':'unlimited','accountScope':'BANK1'}})
        self.assertEqual(len(r['matched']),1)

    def test_cashflow_unlimited_means_latest_imported_date(self):
        from test_cashflow import ledger,transaction,opening
        r=calculate_cashflow({'datasets':[ledger([transaction('2026-01-01','A','10')])],'openings':[opening()],'options':{'endDateMode':'unlimited','endDate':'1900-01-01'}})
        self.assertEqual(r['options']['endDate'],'2026-01-01')
        self.assertEqual(r['summaries'][0]['closing'],11000)


if __name__=='__main__':unittest.main()
