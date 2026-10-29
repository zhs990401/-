from collections import Counter
from io import BytesIO
import unittest

import openpyxl
from reconcile import reconcile_projects
from server import export_report
from test_extended import ds, row


def self_report(amounts, **options):
    return reconcile_projects({'datasets': [ds([row(v, other='A') for v in amounts], project='A')],
                               'options': {'oneToMany': True, **options}})


class SelfProjectTests(unittest.TestCase):
    def test_self_debits_and_credits_match_without_duplicate_sources(self):
        for amounts, match_type in [([100, -100], 'one_to_one'), ([100, -60, -40], 'one_to_many'),
                                    ([100, 100, -100, -100], 'amount_count'),
                                    ([100, 100, -60, -40, -60, -40], 'repeated_split')]:
            report = self_report(amounts)
            self.assertEqual(report['matched'][0]['matchType'], match_type)
            self.assertTrue(report['matched'][0]['selfReconciliation'])
            self.assertFalse(report['differences'] or report['ambiguous'] or report['invalid'])
            self.assertEqual(report['unitOverview'][0]['matchedCount'], len(amounts))
            self.assertEqual(report['unitOverview'][0]['delta'], 0)
            observed = Counter()
            for e in report['matched']:
                for side in ('left', 'right'):
                    observed.update(r['row'] for r in e.get(side+'Rows') or [e[side]])
            self.assertEqual(observed, Counter(range(2, len(amounts)+2)))
            self.assertFalse(any('补齐' in w for w in report['warnings']))

    def test_single_row_never_matches_itself_and_net_is_counted_once(self):
        for amounts in ([40], [-40], [100, -60]):
            report = self_report(amounts)
            self.assertFalse(report['matched'])
            self.assertEqual(report['unitOverview'][0]['delta'], sum(amounts)*100)
            self.assertEqual(report['unitOverview'][0]['reciprocalAmount'], 0)
            self.assertEqual(report['unitOverview'][0]['pendingCount'], len(amounts))
            self.assertEqual(report['projectPairs'][0]['delta'], sum(amounts)*100)

    def test_other_projects_and_currencies_do_not_mix(self):
        datasets = [ds([row(100, other='A'), row(-100, other='B'), row(-100, other='A', currency='USD')], project='A'),
                    ds([row(100, other='B')], project='B')]
        report = reconcile_projects({'datasets': datasets, 'options': {'oneToMany': True}})
        self.assertFalse(report['matched'])
        self.assertEqual(len(report['differences']), 4)

    def test_alias_dates_reference_and_period_are_respected(self):
        a = row(100, 'X', other='Old A')
        b = row(-100, 'Y', other='A', dt='2026-03-01')
        payload = {'datasets': [ds([a, b], project='A')], 'aliases': {'Old A': 'A'}}
        for options in ({'days': 0}, {'days': 'unlimited', 'strictReference': True}):
            report = reconcile_projects({**payload, 'options': options})
            self.assertFalse(report['matched'])
        report = reconcile_projects({**payload, 'options': {'days': 'unlimited'}})
        self.assertEqual(len(report['matched']), 1)
        report = reconcile_projects({**payload, 'options': {'endDate': '2026-02-01'}})
        self.assertEqual(len(report['excludedPeriod']), 1)
        self.assertEqual(report['unitOverview'][0]['delta'], 10000)

    def test_self_summary_and_all_members_are_exported(self):
        report = self_report([100, -60, -40])
        wb = openpyxl.load_workbook(BytesIO(export_report(report)))
        ws = wb['本项目自挂往来汇总']
        self.assertEqual(ws.max_row, 2)
        self.assertEqual([ws.cell(2, n).value for n in (1, 3, 4, 5, 6, 7)], ['A', '1', '2', '100.00', '100.00', '0.00'])
        self.assertIn('本项目自挂往来', wb['已匹配'].cell(2, 3).value)
        wb.close()
