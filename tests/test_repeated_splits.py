"""Fictional ledgers exercising repeated split patterns, not real project data."""
from collections import Counter
from io import BytesIO
import unittest

import openpyxl
from test_extended import bank, ds, row
from reconcile import reconcile, reconcile_projects
from server import export_report


class RepeatedSplitTests(unittest.TestCase):
    def test_repeated_total_and_parts_are_one_statistical_group(self):
        a = [row(100, dt='2026-01-01'), row(100, dt='2026-03-01')]
        b = [row(60), row(40), row(40, dt='2026-02-01'), row(60, dt='2026-02-01')]
        r = bank(a, b, days='unlimited')
        self.assertFalse(r['ambiguous'] or r['differences'] or r['duplicateAmountChecks'])
        self.assertEqual(len(r['matched']), 1)
        e = r['matched'][0]
        self.assertEqual(e['matchType'], 'repeated_split')
        self.assertTrue(e['statisticalMatch'])
        self.assertEqual((e['leftCount'], e['rightCount']), (2, 4))
        self.assertEqual(e['left']['amount'], e['right']['amount'])

    def test_both_sides_and_both_directions(self):
        for sign in (1, -1):
            for direction in ('same', 'reverse'):
                for swap in (False, True):
                    a, b = [100, 100], [60, 40, 60, 40]
                    if swap:
                        a, b = b, a
                    factor = -1 if direction == 'reverse' else 1
                    r = bank([row(v * sign) for v in a], [row(v * sign * factor) for v in b], direction=direction)
                    self.assertEqual(r['matched'][0]['matchType'], 'repeated_split')
                    self.assertFalse(r['ambiguous'] or r['differences'])

    def test_same_valued_parts_use_multiplicity_per_anchor(self):
        r = bank([row(100)] * 2, [row(50)] * 4)
        self.assertEqual(r['matched'][0]['rightCount'], 4)
        self.assertFalse(r['ambiguous'] or r['duplicateAmountChecks'])
        r = bank([row(100)] * 2, [row(50)] * 3)
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']), 5)

    def test_missing_or_excess_parts_do_not_produce_partial_matches(self):
        for b in ([60, 40, 60], [60, 40, 60, 40, 40]):
            r = bank([row(100)] * 2, [row(v) for v in b])
            self.assertFalse(r['matched'])
            self.assertEqual(len(r['ambiguous']), 2 + len(b))

    def test_distinct_amount_patterns_remain_manual(self):
        r = bank([row(100)] * 2, [row(v) for v in [60, 40, 70, 30]])
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']), 6)

    def test_competing_anchor_cannot_consume_batch_members(self):
        r = bank([row(100), row(100), row(80)], [row(60), row(40), row(60), row(40)])
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']), 7)

    def test_unique_reference_match_is_preserved_before_splits(self):
        r = bank([row(100, 'X'), row(100), row(100)], [row(100, 'X'), row(60), row(40), row(60), row(40)])
        self.assertEqual(Counter(e['matchType'] for e in r['matched']), {'one_to_one': 1, 'repeated_split': 1})
        self.assertFalse(r['ambiguous'] or r['differences'] or r['duplicateAmountChecks'])

    def test_dates_are_checked_for_each_constituent(self):
        a = [row(100, dt='2026-01-02')] * 2
        b = [row(v, dt='2026-01-01') for v in [60, 40, 60, 40]]
        r = bank(a, b, days=1)
        self.assertEqual(r['matched'][0]['matchType'], 'repeated_split')
        r = bank(a, b, days=0)
        self.assertFalse(r['matched'])

    def test_same_date_windows_can_form_independent_batches(self):
        a = [row(100, dt=dt) for dt in ['2026-01-01', '2026-01-01', '2026-02-01', '2026-02-01']]
        b = [row(v, dt=dt) for dt in ['2026-01-01', '2026-02-01'] for v in [60, 40, 60, 40]]
        r = bank(a, b)
        self.assertEqual(len(r['matched']), 2)
        self.assertTrue(all(e['matchType'] == 'repeated_split' for e in r['matched']))
        self.assertFalse(r['ambiguous'] or r['differences'])

    def test_currency_account_reference_and_sign_stay_separate(self):
        for field, value in [(5, 'USD'), (6, 'OTHER'), (3, 'OTHER')]:
            b = [row(v, 'X', account='A') for v in [60, 40, 60, 40]]
            b[-1][field] = value
            r = bank([row(100, 'X', account='A')] * 2, b, matchAccount=True, strictReference=True)
            self.assertFalse(r['matched'])
        r = bank([row(100)] * 2, [row(60), row(40), row(120), row(-20)])
        self.assertFalse(r['matched'])

    def test_switches_and_per_anchor_part_limit_are_respected(self):
        a, b = ds([row(100)] * 2), ds([row(v) for v in [55, 30, 15, 55, 30, 15]])
        for options in ({'oneToMany': False}, {'oneToMany': True, 'matchDuplicateCounts': False}, {'oneToMany': True, 'maxParts': 2}):
            r = reconcile({'left': a, 'right': b, 'options': options})
            self.assertFalse(r['matched'])
        r = reconcile({'left': a, 'right': b, 'options': {'oneToMany': True, 'maxParts': 3}})
        self.assertEqual(r['matched'][0]['rightCount'], 6)

    def test_project_overview_and_source_conservation(self):
        a = ds([row(-100, other='B')] * 2, project='A')
        b = ds([row(v, other='A') for v in [60, 40, 60, 40]], project='B')
        r = reconcile_projects({'datasets': [a, b], 'options': {'oneToMany': True}})
        self.assertEqual(r['matched'][0]['matchType'], 'repeated_split')
        self.assertEqual(sorted(v['matchedCount'] for v in r['unitOverview']), [2, 4])
        self.assertTrue(all(v['pendingCount'] == 0 and v['delta'] == 0 for v in r['unitOverview']))
        observed = Counter((v['project'], v['row']) for e in r['matched'] for s in ('left', 'right') for v in e[s + 'Rows'])
        self.assertEqual(observed, Counter({('A', 2): 1, ('A', 3): 1, ('B', 2): 1, ('B', 3): 1, ('B', 4): 1, ('B', 5): 1}))

    def test_matched_split_does_not_hide_true_same_amount_quantity_difference(self):
        a = [row(100), row(100), row(10, currency='USD')]
        b = [row(60), row(40), row(60), row(40), row(10, currency='USD'), row(10, currency='USD')]
        r = bank(a, b)
        self.assertEqual(len(r['matched']), 1)
        self.assertEqual(len(r['ambiguous']), 3)
        self.assertEqual(len(r['duplicateAmountChecks']), 1)
        self.assertEqual(r['duplicateAmountChecks'][0]['currency'], 'USD')

    def test_excel_preserves_all_members_without_invented_pairings(self):
        r = bank([row(100)] * 2, [row(v) for v in [60, 40, 60, 40]])
        wb = openpyxl.load_workbook(BytesIO(export_report(r)))
        ws = wb['已匹配']
        headers = [c.value for c in ws[1]]
        li, ri = headers.index('左侧文件'), headers.index('右侧文件')
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 6)
        self.assertTrue(all(bool(v[li]) != bool(v[ri]) for v in rows))
        self.assertEqual(len({v[0] for v in rows}), 1)
        self.assertEqual(wb['重复或组合冲突'].max_row, 1)
        self.assertEqual(wb['重复金额数量核对'].max_row, 1)
        wb.close()


if __name__ == '__main__':
    unittest.main()
