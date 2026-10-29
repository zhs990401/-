from io import BytesIO
import unittest
import openpyxl
from test_extended import ds,row,bank
from reconcile import reconcile,reconcile_projects
from server import export_report


class DuplicateCountTests(unittest.TestCase):
    def test_equal_repeats_match_as_group_without_individual_pairs(self):
        r=bank([row(100),row(100)],[row(100),row(100)])
        self.assertFalse(r['ambiguous'] or r['differences'])
        self.assertEqual(len(r['matched']),1)
        e=r['matched'][0]
        self.assertEqual(e['matchType'],'amount_count')
        self.assertTrue(e['statisticalMatch'])
        self.assertEqual((e['leftCount'],e['rightCount']),(2,2))
        self.assertEqual(len(e['leftRows'])+len(e['rightRows']),4)
        self.assertEqual(r['duplicateAmountChecks'][0]['status'],'matched')

    def test_unequal_counts_hold_all_unresolved_rows(self):
        r=bank([row(100)]*3,[row(100)]*2)
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']),5)
        c=r['duplicateAmountChecks'][0]
        self.assertEqual((c['leftCount'],c['rightCount'],c['quantityDelta']),(3,2,1))
        self.assertTrue(all('数量不一致' in e['reason'] for e in r['ambiguous']))

    def test_existing_unique_reference_matches_are_preserved(self):
        r=bank([row(100,'X'),row(100,'Y')],[row(100,'X')])
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(r['matched'][0]['matchType'],'one_to_one')
        self.assertEqual(len(r['ambiguous']),1)
        c=r['duplicateAmountChecks'][0]
        self.assertEqual((c['originalLeftCount'],c['originalRightCount']),(2,1))
        self.assertEqual((c['leftCount'],c['rightCount']),(1,0))

    def test_project_direction_and_currency_scope_remain_correct(self):
        a=ds([row(100,other='B')]*2,project='A')
        b=ds([row(-100,other='A')]*2,project='B')
        r=reconcile_projects({'datasets':[a,b]})
        self.assertEqual(r['matched'][0]['matchType'],'amount_count')
        self.assertTrue(all(v['matchedCount']==2 for v in r['unitOverview']))
        b['rows'][0][5]='USD';b['rows'][1][5]='USD'
        r=reconcile_projects({'datasets':[a,b]})
        self.assertFalse(r['matched'])
        self.assertTrue(all(c['status']=='quantity_mismatch' for c in r['duplicateAmountChecks']))

    def test_accounts_do_not_cancel_each_others_counts(self):
        a=ds([row(100,account='A')]*2)
        b=ds([row(100,account='B')]*2)
        r=reconcile({'left':a,'right':b,'options':{'matchAccount':True}})
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['duplicateAmountChecks']),2)

    def test_date_constraints_are_not_ignored_by_equal_counts(self):
        a=ds([row(100,dt='2020-01-01')]*2)
        b=ds([row(100,dt='2026-01-01')]*2)
        r=reconcile({'left':a,'right':b})
        self.assertFalse(r['matched'])
        self.assertTrue(all(c['status']=='quantity_mismatch' for c in r['duplicateAmountChecks']))
        r=reconcile({'left':a,'right':b,'options':{'days':'unlimited'}})
        self.assertEqual(r['matched'][0]['matchType'],'amount_count')

    def test_date_tolerance_group_has_feasible_pairing(self):
        a=ds([row(100,dt='2026-01-01')]*2)
        b=ds([row(100,dt='2026-01-02')]*2)
        r=reconcile({'left':a,'right':b,'options':{'days':1}})
        self.assertEqual(r['matched'][0]['leftCount'],2)

    def test_equal_counts_in_connected_window_still_require_feasible_dates(self):
        a=ds([row(100,dt='2026-01-01')]*3+[row(100,dt='2026-01-03')])
        b=ds([row(100,dt='2026-01-01'),row(100,dt='2026-01-03')]+[row(100,dt='2026-01-04')]*2)
        r=reconcile({'left':a,'right':b,'options':{'days':2}})
        self.assertFalse(r['matched'])
        self.assertEqual(r['duplicateAmountChecks'][0]['status'],'conditions_not_met')
        self.assertEqual(len(r['ambiguous']),8)

    def test_date_components_match_independent_explicit_graph(self):
        import random
        from datetime import date,timedelta
        from duplicate_counts import date_components
        rng=random.Random(1904)
        for _ in range(100):
            d=rng.randint(1,3)
            a=[{'date':(date(2026,1,1)+timedelta(days=rng.randrange(10))).isoformat()} for _ in range(rng.randrange(7))]
            b=[{'date':(date(2026,1,1)+timedelta(days=rng.randrange(10))).isoformat()} for _ in range(rng.randrange(7))]
            nodes=[('L',i) for i in range(len(a))]+[('R',j) for j in range(len(b))]
            adjacency={node:set() for node in nodes}
            for u in nodes:
                for v in nodes:
                    du=(a if u[0]=='L' else b)[u[1]]['date']
                    dv=(a if v[0]=='L' else b)[v[1]]['date']
                    if (u[0]==v[0] and du==dv) or (u[0]!=v[0] and abs((date.fromisoformat(du)-date.fromisoformat(dv)).days)<=d):
                        adjacency[u].add(v)
            expected=set();unseen=set(nodes)
            while unseen:
                stack=[unseen.pop()];component=set(stack)
                while stack:
                    node=stack.pop()
                    for neighbor in adjacency[node]&unseen:
                        unseen.remove(neighbor);component.add(neighbor);stack.append(neighbor)
                expected.add(frozenset(component))
            actual={frozenset([('L',i) for i in il]+[('R',j) for j in ir]) for il,ir in date_components(a,b,list(range(len(a))),list(range(len(b))),d)}
            self.assertEqual(actual,expected)

    def test_strict_reference_is_required_and_cannot_be_empty(self):
        for left_ref,right_ref in [('X','Y'),('','')]:
            r=bank([row(100,left_ref)]*2,[row(100,right_ref)]*2,strictReference=True)
            self.assertFalse(r['matched'])
            self.assertEqual(len(r['ambiguous']),4)
        r=bank([row(100,'X')]*2,[row(100,'X')]*2,strictReference=True)
        self.assertEqual(r['matched'][0]['matchType'],'amount_count')

    def test_large_group_uses_full_counts_not_capped_candidate_lists(self):
        r=bank([row(100)]*300,[row(100)]*300,days='unlimited')
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(r['matched'][0]['leftCount'],300)
        self.assertEqual(r['duplicateAmountChecks'][0]['rightCount'],300)

    def test_option_off_preserves_original_unique_only_behavior(self):
        r=bank([row(100)]*2,[row(100)]*2,matchDuplicateCounts=False)
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']),4)
        self.assertFalse(r['duplicateAmountChecks'])

    def test_one_to_many_remains_available_without_same_amount_candidates(self):
        a=[row(100,dt='2026-01-01'),row(100,dt='2026-01-05')]
        b=[row(60,dt='2026-01-01'),row(40,dt='2026-01-01'),row(70,dt='2026-01-05'),row(30,dt='2026-01-05')]
        r=bank(a,b)
        self.assertEqual(len(r['matched']),2)
        self.assertTrue(all(e['matchType']=='one_to_many' for e in r['matched']))
        self.assertFalse(r['duplicateAmountChecks'])

    def test_fee_scope_and_base_currency_differences_remain_separate(self):
        fee=row(-5,description='手续费')
        r=bank([fee]*2,[fee]*2,separateFees=True)
        self.assertEqual(r['fees'][0]['feeStatus'],'matched')
        self.assertEqual(r['duplicateAmountChecks'][0]['category'],'疑似手续费')
        from test_original_currency import ledger,row as original
        r=reconcile_projects({'datasets':[ledger('A',[original('B','10',base_debit='72')]*2),ledger('B',[original('A',credit='10',base_credit='73')]*2)]})
        self.assertFalse(r['differences'])
        self.assertEqual(r['translationDifferences'][0]['baseDelta'],-200)
        self.assertEqual(r['translationDifferences'][0]['matchType'],'amount_count')

    def test_excel_does_not_imply_arbitrary_individual_assignments(self):
        r=bank([row(100)]*2,[row(100)]*2)
        wb=openpyxl.load_workbook(BytesIO(export_report(r)))
        ws=wb['已匹配'];headers=[c.value for c in ws[1]]
        left=headers.index('左侧文件');right=headers.index('右侧文件')
        self.assertEqual(ws.max_row,5)
        for cells in ws.iter_rows(min_row=2,values_only=True):
            self.assertNotEqual(bool(cells[left]),bool(cells[right]))
        self.assertEqual(wb['重复金额数量核对'].max_row,2)
        wb.close()


if __name__=='__main__':unittest.main()
