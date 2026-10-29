import base64
from io import BytesIO
import unittest

import openpyxl
from bank_batch import parse_bank_file, reconcile_bank_batch, export_bank_batch, fee_type


def encode(w, name):
    b=BytesIO();w.save(b);w.close()
    return dict(name=name,content=base64.b64encode(b.getvalue()).decode())


def inner_file():
    w=openpyxl.Workbook();s=w.active;s.title='项目三栏账'
    s.append(['项目三栏账']);s.append(['本位币 USD'])
    s.append(['2026','','账簿名称','凭证编号','外币名称','账户专项编号','账户专项名称','账户专项开户行','摘要','借方发生','','','贷方发生','','','余额'])
    s.append(['月','日','','','','','','','','外币','汇率','金额','外币','汇率','金额','方向','外币','汇率','金额'])
    s.append(['01','','','','MVR','12345678','TEST BANK','','期初余额',0,0,0,0,0,0,'借',100,0,7])
    s.append(['01','02','测试项目','内1','MVR','12345678','','','正常收入/PC1234',50,0,3,0,0,0,'借',150,0,10])
    s.append(['01','03','测试项目','内2','MVR','12345678','','','正常支出/PC1235',0,0,0,20,0,2,'借',130,0,8])
    s.append(['01','','','','MVR','12345678','','','本期合计',50,0,3,20,0,2,'借',130,0,8])
    s.append(['02','03','测试项目','内3','MVR','12345678','','','期间外收入',20,0,1,0,0,0,'借',150,0,9])
    return encode(w,'虚构内账.xlsx')


def outer_file():
    w=openpyxl.Workbook();w.remove(w.active)
    for account,currency,adjust in [('12345678','MVR',True),('87654321','USD',False)]:
        s=w.create_sheet('Bank MVR '+account)
        s.append(['Posting Date','Due Date','Series','Doc. No.','Trans. No.','Remarks','Offset Acct','Offset Acct Name','Deb./Cred. (FC) (currency)','Deb./Cred. (FC)','Cumulative Balance (FC)','Ref. 2 (Header)','Project'])
        s.append([]);s.append(['Assets',1,None,None,None,'Bank '+currency+' '+account,None,None,None,None,100])
        s.append(['02/01/26','','','外1',1,'正常收入','001','交易',currency,50,150,'PC1234','A'])
        s.append(['03/01/26','','','外2',2,'正常支出','001','交易',currency,-20,130,'PC1235','A'])
        s.append(['05/01/26','','','外3',3,'Bank charges','611105','Bank charges',currency,-5,125,'OD Interest Jan','A'])
        s.append(['06/01/26','','','外4',4,'Bank charges refund','611105','Bank charges',currency,1,126,'refund','A'])
        s.append([None]*7+['Ledger Balance as of 31/01/2026',None,None,126])
        if adjust:s.append(['31/01/26','','','未达1',5,'未达支出','001','交易',currency,-10,136,'PC9999','A'])
        s.append([None]*7+['Bank Balance as of 31/01/2026',None,None,136 if adjust else 126])
    return encode(w,'Account Summary 01 Jan 2026 - 31 Jan 2026 TEST.xlsx')


def sample_report():
    return reconcile_bank_batch(dict(left=[parse_bank_file(inner_file())],right=[parse_bank_file(outer_file())],startDate='2026-01-01',endDate='2026-01-31'))


class BankBatchTests(unittest.TestCase):
    def test_original_currency_and_period_balances_exclude_base_and_future(self):
        r=sample_report();a=r['accounts'][0]
        self.assertEqual(a['left']['opening'],10000)
        self.assertEqual(a['left']['closing'],13000)
        self.assertTrue(a['left']['balanceCheck'])
        self.assertEqual(a['left']['count'],2)
        self.assertTrue(any(x.get('description')=='期间外收入' for x in r['excluded']))

    def test_overdraft_fees_refunds_and_difference_explanation(self):
        r=sample_report();a=r['accounts'][0]
        self.assertEqual((a['openingDelta'],a['closingDelta'],a['feeNetDelta'],a['adjustedClosingDelta']),(0,400,400,0))
        self.assertEqual(a['status'],'期末差异由手续费解释')
        self.assertEqual((a['right']['feeExpense'],a['right']['feeRefund'],a['right']['feeCount']),(500,100,2))
        self.assertEqual(fee_type('OD Intrest Jan'),'透支利息')
        self.assertEqual(fee_type('FD Interest Income','721005'),'')
        self.assertEqual(fee_type('Visa fees and Rental fees'),'')

    def test_adjustments_are_not_transactions_and_missing_side_never_passes(self):
        r=sample_report();a=r['accounts'][0]
        self.assertEqual(a['right']['closing'],12600)
        self.assertEqual(a['right']['bankClosing'],13600)
        self.assertEqual(len(r['adjustments']),1)
        self.assertEqual(a['right']['count'],4)
        self.assertEqual(r['accounts'][1]['status'],'缺少内账')
        self.assertTrue(any('Sheet名含MVR' in w for w in r['warnings']))

    def test_duplicate_accounts_and_missing_opening_are_not_silently_accepted(self):
        left=parse_bank_file(inner_file());right=parse_bank_file(outer_file())
        payload=dict(left=[left],right=[right],startDate='2026-01-01',endDate='2026-01-31')
        with self.assertRaises(ValueError):reconcile_bank_batch({**payload,'left':[left,left]})
        left['accounts'][0]['opening']=None
        r=reconcile_bank_batch(payload)
        self.assertIsNone(r['accounts'][0]['openingDelta'])
        self.assertEqual(r['accounts'][0]['status'],'余额来源/期间待核实')

    def test_date_rule_and_manual_fee_classification_change_results(self):
        left=parse_bank_file(inner_file());right=parse_bank_file(outer_file())
        left['accounts'][0]['transactions'][0]['date']='2026-01-12'
        payload=dict(left=[left],right=[right],startDate='2026-01-01',endDate='2026-01-31',days=0)
        r=reconcile_bank_batch(payload);self.assertGreater(r['accounts'][0]['pending'],0)
        right['accounts'][0]['transactions'][0]['feeType']='人工确认手续费'
        r=reconcile_bank_batch({**payload,'days':'unlimited'})
        self.assertEqual(r['accounts'][0]['right']['feeCount'],3)

    def test_export_keeps_sources_and_prevents_formula_cells(self):
        r=sample_report();r['fees'][0]['description']='=1+1'
        w=openpyxl.load_workbook(BytesIO(export_bank_batch(r)))
        self.assertEqual(w['账户核对总览'].max_row,3)
        self.assertEqual(w['外账余额调节项目'].max_row,2)
        self.assertFalse(any(c.data_type=='f' for s in w for row in s for c in row))
        w.close()

    def test_unrecognized_old_base_only_format_rejected(self):
        w=openpyxl.Workbook();w.active.append(['辅助三栏账','借方','贷方'])
        with self.assertRaises(ValueError):parse_bank_file(encode(w,'old.xlsx'))

    def test_period_and_bad_source_balance_never_claim_consistency(self):
        left=parse_bank_file(inner_file());right=parse_bank_file(outer_file())
        payload=dict(left=[left],right=[right],startDate='2026-01-01',endDate='2026-01-03')
        r=reconcile_bank_batch(payload)
        self.assertIsNone(r['accounts'][0]['right']['bankClosing'])
        self.assertFalse(r['adjustments'])
        left['accounts'][0]['checkpoints'][0]['balance']=99999
        r=reconcile_bank_batch({**payload,'endDate':'2026-01-31'})
        self.assertEqual(r['accounts'][0]['status'],'余额来源/期间待核实')

    def test_new_public_endpoints_process_and_export_without_login(self):
        from test_hosting import HostingTests
        host=HostingTests();host.setUp()
        left=host.call('/api/bank-batch/parse','POST',inner_file())
        self.assertTrue(left['status'].startswith('200'))
        import json
        payload=dict(left=[json.loads(left['body'])],right=[parse_bank_file(outer_file())],startDate='2026-01-01',endDate='2026-01-31')
        result=host.call('/api/bank-batch/reconcile','POST',payload)
        self.assertTrue(result['status'].startswith('200'))
        result=host.call('/api/bank-batch/export','POST',json.loads(result['body']))
        self.assertTrue(result['status'].startswith('200'))
        self.assertTrue(result['body'].startswith(b'PK'))

    def test_every_selected_transaction_appears_once(self):
        from collections import Counter
        left=parse_bank_file(inner_file());right=parse_bank_file(outer_file())
        r=reconcile_bank_batch(dict(left=[left],right=[right],startDate='2026-01-01',endDate='2026-01-31'))
        key=lambda v,side:(side,v['source'],v['sheet'],v['row'])
        expected=Counter(key(v,side) for side,f in [('left',left),('right',right)] for g in f['accounts'] for v in g['transactions'] if v['date']<='2026-01-31')
        observed=Counter(key(v,v['side']) for v in r['fees'])
        for category in ('matched','differences','ambiguous'):
            for e in r[category]:
                for side in ('left','right'):
                    observed.update(key(v,side) for v in e.get(side+'Rows') or ([e[side]] if e[side] else []))
        self.assertEqual(observed,expected)
