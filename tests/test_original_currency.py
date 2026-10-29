import base64
from io import BytesIO
import unittest
import openpyxl

from project_import import project_template
from reconcile import reconcile_projects
from server import export_report, parse_file


HEADERS=['日期','对方项目','原币币种','原币借方发生额','原币贷方发生额','本币币种','本币借方发生额','本币贷方发生额','凭证号','科目名称']
MAPPING={'mode':'split','date':'日期','counterparty':'对方项目','currency':'原币币种','debit':'原币借方发生额','credit':'原币贷方发生额','baseCurrency':'本币币种','baseDebit':'本币借方发生额','baseCredit':'本币贷方发生额','reference':'凭证号','subjectName':'科目名称'}


def ledger(project, rows):
    return {'name':project+'.xlsx','headers':list(HEADERS),'rows':rows,'mapping':dict(MAPPING)}


def row(other, debit='', credit='', base_debit='', base_credit='', currency='USD', base_currency='CNY', dt='2026-01-01', ref=''):
    return [dt,other,currency,debit,credit,base_currency,base_debit,base_credit,ref,'其他应收/应付']


class OriginalCurrencyTests(unittest.TestCase):
    def test_original_matches_despite_different_base_amounts_and_vouchers(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','1000',base_debit='7200',ref='A1')]),ledger('B',[row('A',credit='1000',base_credit='7300',ref='B99')])]})
        self.assertEqual(len(r['matched']),1)
        self.assertFalse(r['differences'])
        self.assertEqual(len(r['translationDifferences']),1)
        self.assertEqual(r['translationDifferences'][0]['baseDelta'],-10000)
        self.assertEqual(r['projectPairs'][0]['status'],'原币按规则一致')
        self.assertEqual(r['options']['amountBasis'],'original')
        self.assertEqual(r['importInventory'][0]['project'],'A')

    def test_original_difference_is_not_hidden_by_equal_base_amount(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','1000',base_debit='7200',ref='X')]),ledger('B',[row('A',credit='980',base_credit='7200',ref='X')])]})
        self.assertFalse(r['matched'])
        self.assertEqual(r['differences'][0]['delta'],2000)
        self.assertFalse(r['translationDifferences'])

    def test_different_original_currencies_never_cross_match(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','10',currency='USD')]),ledger('B',[row('A',credit='10',currency='EUR')])]})
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['projectPairs']),2)

    def test_currency_names_are_normalized_without_exchange_conversion(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','10',currency='美元')]),ledger('B',[row('A',credit='10',currency='USD')])]})
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(r['matched'][0]['left']['currency'],'USD')

    def test_different_base_currencies_are_not_compared_or_converted(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','10',base_debit='72',base_currency='CNY')]),ledger('B',[row('A',credit='10',base_credit='80',base_currency='HKD')])]})
        self.assertEqual(len(r['matched']),1)
        self.assertFalse(r['translationDifferences'])
        self.assertEqual(r['baseComparisonSkipped'],1)
        self.assertTrue(any('未进行汇率换算' in w for w in r['warnings']))

    def test_signed_original_amounts_are_supported(self):
        headers=['日期','对方项目','原币币种','原币金额']
        mapping={'mode':'signed','date':'日期','counterparty':'对方项目','currency':'原币币种','amount':'原币金额'}
        a={'name':'A.xlsx','headers':headers,'mapping':mapping,'rows':[['2026-01-01','B','USD','10']]}
        b={'name':'B.xlsx','headers':headers,'mapping':mapping,'rows':[['2026-01-01','A','USD','-10']]}
        r=reconcile_projects({'datasets':[a,b]})
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(r['matched'][0]['left']['amount'],1000)

    def test_missing_row_currency_is_invalid_not_assumed_cny(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','10',currency='')]),ledger('B',[row('A',credit='10',currency='CNY')])]})
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['invalid']),1)
        self.assertIn('不能假设人民币',r['invalid'][0]['reason'])

    def test_base_columns_and_balances_cannot_be_used_as_original(self):
        for name in ['本币借方发生额','原币期末余额']:
            d=ledger('A',[row('B','10')]);d['mapping']['debit']=name;d['headers'].append(name);d['originalConfirmed']=True
            with self.assertRaises(ValueError):
                reconcile_projects({'datasets':[d]})

    def test_generic_columns_need_explicit_confirmation_and_currency(self):
        d=ledger('A',[row('B','10')]);d['headers'][3]='借方';d['mapping']['debit']='借方'
        with self.assertRaises(ValueError):
            reconcile_projects({'datasets':[d]})
        d['originalConfirmed']=True;d['mapping'].pop('currency')
        with self.assertRaises(ValueError):
            reconcile_projects({'datasets':[d]})
        d['currencyName']='USD'
        r=reconcile_projects({'datasets':[d]})
        self.assertEqual(r['differences'][0]['left']['currency'],'USD')

    def test_one_to_many_preserves_base_sum_and_all_original_members(self):
        a=ledger('A',[row('B','1000',base_debit='7300')])
        b=ledger('B',[row('A',credit='600',base_credit='4300'),row('A',credit='400',base_credit='2900')])
        r=reconcile_projects({'datasets':[a,b],'options':{'oneToMany':True}})
        self.assertEqual(r['matched'][0]['right']['baseAmount'],-720000)
        self.assertEqual(len(r['translationDifferences'][0]['rightRows']),2)
        self.assertEqual(r['translationDifferences'][0]['baseDelta'],10000)

    def test_invalid_base_data_does_not_invalidate_original_match(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','10',base_debit='oops')]),ledger('B',[row('A',credit='10',base_credit='72')])]})
        self.assertEqual(len(r['matched']),1)
        self.assertFalse(r['translationDifferences'])
        self.assertTrue(any('本币资料' in w for w in r['warnings']))

    def test_period_filter_preserves_excluded_rows(self):
        a=ledger('A',[row('B','10',dt='2025-12-31'),row('B','20',dt='2026-01-01')]);b=ledger('B',[row('A',credit='20')])
        r=reconcile_projects({'datasets':[a,b],'options':{'startDate':'2026-01-01','endDate':'2026-01-31'}})
        self.assertEqual(len(r['matched']),1)
        self.assertEqual(len(r['excludedPeriod']),1)
        self.assertEqual(r['matched'][0]['left']['row'],3)
        with self.assertRaises(ValueError):
            reconcile_projects({'datasets':[a,b],'options':{'startDate':'2026-02-01','endDate':'2026-01-31'}})

    def test_duplicate_project_files_are_rejected(self):
        d=ledger('A',[row('B','10')]);copy={**d,'name':'副本.xlsx','projectName':'A'}
        with self.assertRaises(ValueError):
            reconcile_projects({'datasets':[d,copy]})

    def test_excel_export_has_original_basis_and_separate_base_difference(self):
        r=reconcile_projects({'datasets':[ledger('A',[row('B','10',base_debit='72')]),ledger('B',[row('A',credit='10',base_credit='73')])]})
        wb=openpyxl.load_workbook(BytesIO(export_report(r)),data_only=False)
        self.assertEqual(wb['已匹配'].max_row,2)
        self.assertEqual(wb['本币金额差异'].max_row,2)
        self.assertEqual(wb['项目导入检查'].max_row,3)
        self.assertTrue(any(c.value=='-1.00' for row_ in wb['本币金额差异'] for c in row_))
        self.assertFalse(any(c.data_type=='f' for ws in wb for row_ in ws for c in row_))
        wb.close()

    def test_auto_header_and_detail_sheet_detection(self):
        wb=openpyxl.Workbook();wb.active.title='说明';wb.active.append(['请使用往来明细'])
        ws=wb.create_sheet('往来明细');ws.append(['项目A往来账']);ws.append(['导出期间2026年1月']);ws.append(HEADERS);ws.append(row('B','10'))
        stream=BytesIO();wb.save(stream);wb.close()
        r=parse_file({'name':'A.xlsx','content':base64.b64encode(stream.getvalue()).decode(),'autoHeader':True})
        self.assertEqual(r['sheet'],'往来明细')
        self.assertEqual(r['firstRow'],4)
        self.assertIn('原币借方发生额',r['headers'])

    def test_template_has_required_original_currency_headers(self):
        wb=openpyxl.load_workbook(BytesIO(project_template()))
        self.assertEqual(wb.sheetnames[0],'往来明细')
        self.assertIn('原币币种',[c.value for c in wb['往来明细'][1]])
        self.assertIn('原币借方发生额',[c.value for c in wb['往来明细'][1]])
        wb.close()


if __name__=='__main__':
    unittest.main()
