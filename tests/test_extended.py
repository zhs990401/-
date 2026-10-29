import unittest
from io import BytesIO
from unittest.mock import patch

import openpyxl
from reconcile import reconcile, reconcile_projects
from server import export_report


HEADERS = ['日期', '借方', '贷方', '参考号', '摘要', '币种', '账户', '本方项目', '对方项目']
MAPPING = dict(mode='split', date='日期', debit='借方', credit='贷方', reference='参考号', description='摘要', currency='币种', account='账户')


def ds(rows, name='账簿.xlsx', project=None, own_column=False):
    mapping = {**MAPPING}
    if project or own_column:
        mapping['counterparty'] = '对方项目'
    if own_column:
        mapping['project'] = '本方项目'
    return dict(headers=HEADERS, rows=rows, mapping=mapping, name=name, projectName=project or '',originalConfirmed=True)


def row(amount, reference='', description='', currency='CNY', account='', own='', other='', dt='2026-01-01'):
    return [dt, str(amount) if amount > 0 else '', str(-amount) if amount < 0 else '', reference, description, currency, account, own, other]


def bank(left, right, **options):
    return reconcile(dict(left=ds(left, '内账.xlsx'), right=ds(right, '外账.xlsx'), options=dict(oneToMany=True, **options)))


class GroupTests(unittest.TestCase):
    def test_every_input_row_is_reported_once_including_fee_groups(self):
        import random
        from collections import Counter
        rng = random.Random(712)
        for _ in range(80):
            a = [row(rng.choice([-1, 1]) * rng.randint(1, 9), description=rng.choice(['转账', '手续费'])) for _ in range(6)]
            b = [row(rng.choice([-1, 1]) * rng.randint(1, 9), description=rng.choice(['转账', '手续费'])) for _ in range(6)]
            report = bank(a, b, separateFees=True)
            observed = Counter()
            for category in ('matched', 'differences', 'ambiguous', 'invalid', 'fees'):
                for entry in report[category]:
                    for side in ('left', 'right'):
                        members = entry.get(side + 'Rows') or ([entry[side]] if entry[side] else [])
                        observed.update((side, r['row']) for r in members)
            self.assertEqual(observed, Counter({(side, line): 1 for side in ('left', 'right') for line in range(2, 8)}))

    def test_unique_one_to_many_both_directions_and_origin(self):
        for a, b in (([row(-100)], [row(-60), row(-40)]), ([row(60), row(40)], [row(100)])):
            r = bank(a, b)
            self.assertEqual(len(r['matched']), 1)
            match = r['matched'][0]
            self.assertEqual(match['matchType'], 'one_to_many')
            self.assertEqual(len(match['leftRows']) + len(match['rightRows']), 3)
            self.assertEqual(match['delta'], 0)
            self.assertTrue(all(v['source'] for v in match['leftRows'] + match['rightRows']))
            self.assertFalse(r['differences'] or r['ambiguous'])

    def test_multiple_valid_subsets_are_not_chosen(self):
        r = bank([row(100)], [row(60), row(40), row(70), row(30)])
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']), 5)

    def test_shared_constituents_block_all_competing_groups(self):
        r = bank([row(100), row(100)], [row(60), row(40)])
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['ambiguous']), 4)

    def test_does_not_net_opposite_directions(self):
        r = bank([row(100)], [row(150), row(-50)])
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['differences']), 3)

    def test_same_currency_account_and_strict_reference_required(self):
        for pieces, opts in (([row(60, 'X'), row(40, 'Y')], dict(strictReference=True)),
                              ([row(60), row(40, currency='USD')], {}),
                              ([row(60, account='A'), row(40, account='B')], dict(matchAccount=True))):
            r = bank([row(100, 'X', account='A')], pieces, **opts)
            self.assertFalse(r['matched'])

    def test_reference_optional_and_reverse_group(self):
        r = bank([row(100, 'A')], [row(-60, 'A'), row(-40, 'B')], direction='reverse')
        self.assertEqual(len(r['matched']), 1)

    def test_constituents_limited_and_search_failure_is_explicit(self):
        r = bank([row(100)], [row(20)] * 5, maxParts=4)
        self.assertFalse(r['matched'])
        r = bank([row(100)], [row(5)] * 20)
        self.assertFalse(r['matched'])
        self.assertTrue(r['warnings'])
        self.assertEqual(len(r['ambiguous']), 21)
        with patch('reconcile.GROUP_BUDGET', 1):
            r = bank([row(100)], [row(60), row(40), row(20)])
        self.assertFalse(r['matched'])
        self.assertIn('计算上限', r['warnings'][0])

    def test_one_to_one_precedes_group_search(self):
        r = bank([row(100)], [row(100), row(60), row(40)])
        self.assertEqual(len(r['matched']), 1)
        self.assertEqual(r['matched'][0]['matchType'], 'one_to_one')
        self.assertEqual(len(r['differences']), 2)


class FeeTests(unittest.TestCase):
    def test_missing_fees_and_refunds_separate_without_hiding_delta(self):
        r = bank([row(-100)], [row(-100), row(-12, description='银行手续费'), row(2, description='手续费退回')], separateFees=True)
        self.assertEqual(len(r['matched']), 1)
        self.assertEqual(len(r['fees']), 2)
        self.assertEqual(r['totals'][0]['delta'], 1000)
        s = r['feeSummary'][0]
        self.assertEqual(s['missingLeftExpense'], 1200)
        self.assertEqual(s['missingLeftRefund'], 200)
        self.assertEqual(s['rightCount'], 2)
        self.assertTrue(all(e['feeStatus'] == 'differences' for e in r['fees']))

    def test_fees_present_on_both_sides_and_duplicates(self):
        fee = row(-5, description='手续费')
        r = bank([fee], [fee], separateFees=True)
        self.assertEqual(r['fees'][0]['feeStatus'], 'matched')
        self.assertEqual(r['feeSummary'][0]['missingLeftExpense'], 0)
        r = bank([fee, fee], [fee], separateFees=True)
        self.assertTrue(all(e['feeStatus'] == 'ambiguous' for e in r['fees']))
        self.assertEqual(r['feeSummary'][0]['missingLeftCount'], 0)

    def test_fee_currency_keywords_and_no_cross_category_combination(self):
        r = bank([row(-100)], [row(-90), row(-10, description='手续费'), row(-3, description='手续费', currency='USD')], separateFees=True)
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['feeSummary']), 2)
        r = bank([], [row(-2, description='电报费')], separateFees=True, feeKeywords='电报费')
        self.assertEqual(len(r['fees']), 1)


class ProjectTests(unittest.TestCase):
    def test_three_projects_are_matched_only_with_reciprocal_counterparty(self):
        datasets = [ds([row(100, other='B'), row(50, other='C')], project='A'),
                    ds([row(-60, other='A'), row(-40, other='A'), row(20, other='C')], project='B'),
                    ds([row(-49, other='A'), row(-20, other='B')], project='C')]
        r = reconcile_projects(dict(datasets=datasets, options=dict(oneToMany=True)))
        self.assertEqual(len(r['matched']), 2)
        self.assertEqual(len(r['differences']), 2)
        self.assertEqual(len(r['projectPairs']), 3)
        ac = next(p for p in r['projectPairs'] if p['projectA'] == 'A' and p['projectB'] == 'C')
        self.assertEqual(ac['delta'], 100)
        self.assertTrue(all(e['left']['project'] == e['right']['counterparty'] and e['right']['project'] == e['left']['counterparty'] for e in r['matched']))

    def test_same_amount_between_different_projects_never_cross_matches(self):
        r = reconcile_projects(dict(datasets=[ds([row(100, other='B')], project='A'), ds([row(-100, other='C')], project='B'), ds([], project='C')]))
        self.assertFalse(r['matched'])
        self.assertEqual(len(r['differences']), 2)

    def test_missing_project_and_zero_net_still_pending(self):
        r = reconcile_projects(dict(datasets=[ds([row(100, other='D'), row(-100, other='D')], project='A')]))
        self.assertEqual(r['projectPairs'][0]['delta'], 0)
        self.assertEqual(r['projectPairs'][0]['status'], '原币待核实')
        self.assertEqual(r['projectPairs'][0]['missingProject'], 'D')
        self.assertTrue(r['warnings'])

    def test_combined_sheet_aliases_and_invalid_project_rows(self):
        rows = [row(100, own='甲', other='乙'), row(-100, own='B', other='A'), row(2, own='A', other='A'), row(3, own='A', other='')]
        r = reconcile_projects(dict(datasets=[ds(rows, own_column=True)], aliases={'甲':'A','乙':'B'}))
        self.assertEqual(len(r['matched']), 1)
        self.assertEqual(len(r['invalid']), 2)
        self.assertEqual(r['matched'][0]['left']['originalProject'], '甲')
        self.assertEqual(r['counts']['projects'], 2)

    def test_missing_mapping_and_alias_cycle_are_rejected(self):
        for payload in (dict(datasets=[ds([row(100)])]), dict(datasets=[ds([], project='A')], aliases={'A':'B','B':'A'})):
            with self.assertRaises(ValueError):
                reconcile_projects(payload)


class ExportTests(unittest.TestCase):
    def test_workbook_preserves_group_members_and_safe_text(self):
        r = bank([row(-100, description='=HYPERLINK("bad")')], [row(-60), row(-40), row(-3, description='手续费')], separateFees=True)
        content = export_report(r)
        wb = openpyxl.load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(wb['已匹配'].max_row, 3)
        self.assertEqual(wb['手续费汇总'].max_row, 2)
        cells = [c for sheet in wb for row_ in sheet for c in row_]
        self.assertTrue(any(c.value == '=HYPERLINK("bad")' for c in cells))
        self.assertFalse(any(c.data_type == 'f' for c in cells))
        self.assertEqual(wb['疑似手续费明细'].max_row, 2)
        wb.close()


if __name__ == '__main__':
    unittest.main()
