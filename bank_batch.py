"""Account/currency isolated bank books, including original-currency balances."""
import base64
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import re
from zipfile import ZipFile

import openpyxl
from project_import import currency_name
from reconcile import compare_records, settings, safe_total, money


def cents(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
        rounded = number.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        if abs(number-rounded) > Decimal('.000001'):
            raise ValueError('原币金额超过两位小数，请检查导出格式')
        return money(str(rounded))
    return money(value)


def fee_type(text, code=''):
    if re.search(r'透支利息|overdraft.*int(?:erest|rest)|\bOD\b.*int(?:erest|rest)|int(?:erest|rest).*\bOD\b', text, re.I) or (str(code)=='611105' and re.search(r'利息|interest|intrest',text,re.I)):
        return '透支利息'
    if str(code) == '611105' or re.search(r'手续费|汇划费|结算费|账户管理费|网银服务费|银行工本费|\bbank\s*(?:charges?|fees?)\b|\bOD\s+Doc\s+Chgs\b', text, re.I):
        return '银行手续费'
    return ''


def parse_bank_file(payload):
    name = str(payload.get('name', ''))
    if not name.lower().endswith('.xlsx'):
        raise ValueError('批量银行核对目前支持 .xlsx；请导出含原币发生额和余额的 Excel')
    content = base64.b64decode(payload.get('content', ''), validate=True)
    if len(content) > 10*1024*1024:
        raise ValueError('每份文件最多10 MB')
    with ZipFile(BytesIO(content)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100*1024*1024:
            raise ValueError('Excel 解压后过大')
    w = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    groups = {}
    excluded, warnings = [], []
    def group(account, currency, sheet):
        if not account or not currency:
            raise ValueError(f'{sheet}：账号或原币币种缺失，不能默认币种')
        key = (account, currency)
        if key in groups and groups[key]['sheet'] != sheet:
            raise ValueError('同一文件多个Sheet重复出现同账号币种，请核实重复导出或合并期间')
        if key not in groups:
            groups[key] = dict(account=account, currency=currency, source=name, sheet=sheet,
                               opening=None, openingDate='', transactions=[], adjustments=[], checkpoints=[])
        return groups[key]
    def record(g, number, dt, amount, balance, description, reference, **extra):
        return dict(source=name, sheet=g['sheet'], row=number, date=dt, amount=amount,
                    balance=balance, description=description, reference=reference,
                    account=g['account'], currency=g['currency'], feeType=fee_type(description), **extra)
    total_rows = 0
    try:
        if len(w.sheetnames) > 50:
            raise ValueError('每份文件最多50个工作表')
        for ws in w:
            if ws.max_row > 60050 or ws.max_column > 100:
                raise ValueError('工作表超过60,000行或100列')
            rows = list(ws.values)
            total_rows += len(rows)
            if total_rows > 60050:
                raise ValueError('每份文件最多60,000行')
            get = lambda r,i: str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else ''
            if rows and get(rows[0],0) == 'Posting Date':
                if len(rows)<3:
                    raise ValueError('外账缺少期初账户行')
                account_text=get(rows[2],5)
                found=re.findall(r'\d{6,}',account_text)
                if len(found)!=1:
                    raise ValueError(f'{ws.title}：不能从期初账户信息识别唯一账号')
                account=found[0]
                period=re.search(r'(\d{2} [A-Za-z]{3} \d{4})\s*-\s*(\d{2} [A-Za-z]{3} \d{4})',name)
                start,end=(tuple(datetime.strptime(x,'%d %b %Y').date().isoformat() for x in period.groups()) if period else ('',''))
                stage='transactions';active=None
                for number,r in enumerate(rows[3:],4):
                    label=get(r,7)
                    if label.startswith('Ledger Balance as of') or label.startswith('Bank Balance as of'):
                        if active:
                            field='ledgerClosing' if label.startswith('Ledger') else 'bankClosing'
                            active[field]=cents(r[10]);active[field+'Row']=number
                        stage='adjustments' if label.startswith('Ledger') else 'ended'
                        continue
                    if not get(r,0) and not get(r,9):
                        continue
                    try:
                        dt=datetime.strptime(get(r,0),'%d/%m/%y').date().isoformat()
                    except ValueError:
                        raise ValueError(f'{ws.title} 第{number}行交易日期无效') from None
                    currency=currency_name(get(r,8))
                    if active and active['currency'] != currency:
                        raise ValueError(f'{ws.title}：一个账户Sheet出现多币种，无法分配期初余额，请分币种导出')
                    g=group(account,currency,ws.title);active=g
                    if not g['openingDate']:
                        g.update(opening=cents(rows[2][10]),openingDate=start,coverageEnd=end)
                    amount=cents(r[9]);balance=cents(r[10])
                    if amount is None:
                        raise ValueError(f'{ws.title} 第{number}行发生额缺失，可能缺少公式缓存值')
                    description=' / '.join(get(r,i) for i in (5,7,11) if get(r,i))
                    entry=record(g,number,dt,amount,balance,description,get(r,11),voucher=get(r,3))
                    entry['feeType']=fee_type(description,get(r,6))
                    if stage=='ended':
                        raise ValueError('银行期末余额之后仍有交易，请检查工作表结构')
                    g[stage].append(entry)
                    if 'MVR' in ws.title.upper() and currency=='USD':
                        warnings.append(f'{ws.title}：Sheet名含MVR，但金额列为USD；以金额列USD核对')
                if not period:
                    warnings.append(f'{ws.title}：文件名未识别报表期间，请重新导出带起止日期的Account Summary')
            else:
                header=next((i for i,r in enumerate(rows[:50]) if '账户专项编号' in r and '借方发生' in r and '贷方发生' in r and '余额' in r),None)
                if header is None:
                    if any(any(v is not None for v in r) for r in rows):
                        raise ValueError(f'{ws.title}：不支持的银行明细格式，请上传项目三栏账或Account Summary')
                    continue
                parent=[str(v or '').strip() for v in rows[header]]
                child=[str(v or '').strip() for v in rows[header+1]]
                col=lambda label: parent.index(label) if label in parent else None
                ai,ci,di,vi=col('账户专项编号'),col('外币名称'),col('摘要'),col('凭证编号')
                if ci is None:
                    raise ValueError('内账缺少原币币种列')
                year=next((int(v) for v in parent if re.fullmatch(r'20\d{2}',v)),None)
                if not year:
                    raise ValueError('内账未识别报表年份')
                ds,cs,bs=col('借方发生'),col('贷方发生'),col('余额')
                def original(a,b):
                    return next((i for i in range(a,b) if child[i] in ('外币','原币')),None)
                debit,credit,bi=original(ds,cs),original(cs,bs),original(bs,len(child))
                direction=next((i for i in range(bs,len(child)) if child[i]=='方向'),None)
                if any(i is None for i in (debit,credit,bi,direction,vi,di)):
                    raise ValueError('内账原币借贷、余额或方向列不完整')
                for number,r in enumerate(rows[header+2:],header+3):
                    if not get(r,ai):
                        if any(get(r,i) for i in (vi,ai)):
                            raise ValueError(f'第{number}行缺少账号')
                        if any(v is not None for v in r):excluded.append(dict(source=name,sheet=ws.title,row=number,reason='非账户明细/汇总行'))
                        continue
                    g=group(get(r,ai),currency_name(get(r,ci)),ws.title)
                    balance=cents(r[bi]);sign=get(r,direction)
                    if balance is not None:
                        if sign=='贷':balance=-balance
                        elif sign not in ('借','平'):
                            raise ValueError(f'第{number}行余额方向无效')
                    description=get(r,di)
                    if description in ('期初余额','上年结转') and not get(r,vi):
                        if g['opening'] is not None:
                            raise ValueError('同账户币种出现多个期初，请勿重复上传或混合多个期间')
                        g.update(opening=balance,openingDate=date(year,int(float(get(r,0) or 1)),1).isoformat())
                        continue
                    if not get(r,vi):
                        if description in ('本期合计','本月合计'):
                            import calendar
                            m=int(float(get(r,0)));g['checkpoints'].append(dict(date=date(year,m,calendar.monthrange(year,m)[1]).isoformat(),balance=balance,row=number))
                        excluded.append(dict(source=name,sheet=ws.title,row=number,account=g['account'],currency=g['currency'],reason=description or '非逐笔行'))
                        continue
                    dt=date(year,int(float(get(r,0))),int(float(get(r,1)))).isoformat()
                    d,c=cents(r[debit]),cents(r[credit])
                    if d is None or c is None or (d and c):
                        raise ValueError(f'第{number}行原币借贷缺失或同时有金额')
                    if not d and not c:
                        excluded.append(dict(source=name,sheet=ws.title,row=number,account=g['account'],currency=g['currency'],reason='零原币发生，不纳入交易匹配'))
                        continue
                    g['transactions'].append(record(g,number,dt,d-c,balance,description,get(r,vi),project=get(r,col('账簿名称'))))
                for g in groups.values():
                    if g['transactions'] and 'coverageEnd' not in g:
                        g['coverageEnd']=max(r['date'] for r in g['transactions']+g['checkpoints'])
        if not groups:
            raise ValueError('未识别到银行账户')
        return dict(name=name,accounts=list(groups.values()),excluded=excluded,warnings=sorted(set(warnings)))
    finally:
        w.close()


def reconcile_bank_batch(payload):
    start=date.fromisoformat(payload['startDate']).isoformat();end=date.fromisoformat(payload['endDate']).isoformat()
    if start>end:raise ValueError('开始日期不能晚于结束日期')
    report=dict(startDate=start,endDate=end,accounts=[],matched=[],differences=[],ambiguous=[],fees=[],adjustments=[],excluded=[],warnings=[])
    indices={}
    for side in ('left','right'):
        files=payload.get(side,[])
        if not files or len(files)>20:raise ValueError('内外账各需1至20份文件')
        index={};count=0
        for f in files:
            report['excluded'].extend(f.get('excluded',[]));report['warnings'].extend(f.get('warnings',[]))
            for g in f['accounts']:
                key=(g['account'],g['currency'])
                if key in index:raise ValueError('同侧同账号币种重复：'+key[0]+'；请合并期间后上传，避免重复计算')
                count+=len(g['transactions'])+len(g['adjustments'])
                if count>60000:raise ValueError('每侧最多60,000条')
                index[key]=g
        indices[side]=index
    options=settings(dict(days=payload.get('days','unlimited'),oneToMany=True,maxParts=5,matchAccount=True))
    for key in sorted(indices['left'].keys()|indices['right'].keys()):
        account,currency=key
        summary=dict(account=account,currency=currency,left=None,right=None,status='待核实')
        normal={};fees={}
        for side in ('left','right'):
            g=indices[side].get(key)
            normal[side]=[];fees[side]=[]
            if g is None:continue
            tx=g['transactions'];safe_total(tx)
            if any(r['account']!=account or r['currency']!=currency or not isinstance(r['amount'],int) for r in tx):
                raise ValueError('账户明细格式无效')
            selected=[{**r,'side':side} for r in tx if start<=r['date']<=end]
            normal[side]=[r for r in selected if not r['feeType']];fees[side]=[r for r in selected if r['feeType']]
            report['fees'].extend(fees[side])
            if end==g.get('coverageEnd'):
                report['adjustments'].extend({**r,'side':side} for r in g['adjustments'])
            else:
                report['excluded'].extend({**r,'side':side,'reason':'调节项目对应报表截止日，与所选截止日不同'} for r in g['adjustments'])
            report['excluded'].extend({**r,'side':side,'reason':'期间外交易'} for r in tx if not start<=r['date']<=end)
            usable=g['opening'] is not None and bool(g['openingDate']) and g['openingDate']<=start and end<=g.get('coverageEnd','')
            opening=g['opening']+sum(r['amount'] for r in tx if g['openingDate']<=r['date']<start) if usable else None
            closing=opening+sum(r['amount'] for r in selected) if usable else None
            # Source balances independently verify the computed period closing.
            checkpoints=[r for r in g.get('checkpoints',[]) if r['date']==end]
            if checkpoints:source_closing=checkpoints[-1]['balance']
            elif end==g.get('coverageEnd') and g.get('ledgerClosing') is not None:source_closing=g['ledgerClosing']
            else:
                candidates=[r for r in tx if r['date']<=end and r['balance'] is not None]
                source_closing=max(enumerate(candidates),key=lambda v:(v[1]['date'],v[0]))[1]['balance'] if candidates else g['opening']
            fnet=sum(r['amount'] for r in fees[side]);normalnet=sum(r['amount'] for r in normal[side])
            summary[side]=dict(opening=opening,closing=closing,sourceClosing=source_closing,
                               balanceCheck=closing is not None and source_closing is not None and closing==source_closing,
                               adjustedClosing=closing-fnet if closing is not None else None,normalNet=normalnet,feeNet=fnet,
                               feeExpense=sum(-r['amount'] for r in fees[side] if r['amount']<0),feeRefund=sum(r['amount'] for r in fees[side] if r['amount']>0),
                               feeCount=len(fees[side]),count=len(selected),bankClosing=g.get('bankClosing') if end==g.get('coverageEnd') else None,source=g['source'],sheet=g['sheet'])
            if not usable:report['warnings'].append(f'{account}/{currency}：'+('内账' if side=='left' else '外账')+'期初、报表期间或覆盖范围不足，余额不作一致判断')
        a,b=summary['left'],summary['right']
        for field in ('opening','closing','adjustedClosing','feeNet','normalNet'):
            summary[field+'Delta']=a[field]-b[field] if a and b and a[field] is not None and b[field] is not None else None
        comparison=compare_records(normal['left'],normal['right'],options)
        for category in ('matched','differences','ambiguous'):
            report[category].extend({**e,'account':account,'currency':currency} for e in comparison[category])
        report['warnings'].extend(comparison['warnings'])
        summary['pending']=len(comparison['differences'])+len(comparison['ambiguous'])
        if not a or not b:summary['status']='缺少内账' if not a else '缺少外账'
        elif not a['balanceCheck'] or not b['balanceCheck']:summary['status']='余额来源/期间待核实'
        elif summary['openingDelta']==0 and summary['closingDelta']==0 and not summary['pending'] and summary['feeNetDelta']==0:summary['status']='按当前规则一致'
        elif summary['openingDelta']==0 and summary['normalNetDelta']==0 and not summary['pending'] and summary['closingDelta']==summary['feeNetDelta']:summary['status']='期末差异由手续费解释'
        report['accounts'].append(summary)
    report['warnings']=sorted(set(report['warnings']))
    return report


def export_bank_batch(report):
    w=openpyxl.Workbook();w.remove(w.active)
    def add(name,headers,rows):
        s=w.create_sheet(name);s.append(headers)
        for n,row in enumerate(rows):
            if n>=120000:raise ValueError('导出记录过多')
            s.append(['' if v is None else str(v) for v in row])
        for row in s:
            for c in row:c.data_type='s'
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    fmt=lambda n: None if n is None else ('-' if n<0 else '')+str(abs(n)//100)+'.'+str(abs(n)%100).zfill(2)
    add('账户核对总览',['账号','币种','状态','期初差','期末差','剔费期末差','手续费净额差','普通交易净额差','待核实项'],
        ([r['account'],r['currency'],r['status']]+[fmt(r[k+'Delta']) for k in ('opening','closing','adjustedClosing','feeNet','normalNet')]+[r['pending']] for r in report['accounts']))
    add('余额与手续费核对',['账号','币种','内外账','期初','计算期末','来源期末','余额勾稽','剔费期末','手续费支出','手续费退款','手续费净额','手续费笔数','调节后银行余额','文件','Sheet'],
        ([r['account'],r['currency'],side]+[fmt(v[k]) for k in ('opening','closing','sourceClosing')]+[v['balanceCheck']]+[fmt(v[k]) for k in ('adjustedClosing','feeExpense','feeRefund','feeNet')]+[v['feeCount'],fmt(v['bankClosing']),v['source'],v['sheet']] for r in report['accounts'] for side in ('left','right') if (v:=r[side])))
    def recordrows(records):
        for r in records:
            yield [r.get(k,'') for k in ('account','currency','side','source','sheet','row','date','description','reference','feeType','classification')]+[fmt(r.get('amount')),r.get('reason','')]
    headers=['账号','币种','内外账','文件','Sheet','行号','日期','摘要','业务号','费用类型','分类调整','发生额','说明']
    for key,label in [('fees','手续费含透支利息'),('adjustments','外账余额调节项目'),('excluded','未纳入及期间外')]:add(label,headers,recordrows(report[key]))
    for key,label in [('matched','普通交易已匹配'),('differences','普通交易待核实'),('ambiguous','重复与组合冲突')]:
        rows=[]
        for n,e in enumerate(report[key],1):
            for side in ('left','right'):
                members=e.get(side+'Rows') or ([e[side]] if e[side] else [])
                for r in members:rows.append([n,e['reason']]+next(recordrows([{**r,'side':side}])))
        add(label,['组号','匹配说明']+headers,rows)
    add('规则与提示',['项目','内容'],[['核对期间',report['startDate']+' 至 '+report['endDate']],['手续费范围','含透支利息；普通交易配对排除手续费；保留原始余额'],['差异公式','期末差 = 期初差 + 普通交易净额差 + 手续费净额差']]+[['提示',v] for v in report['warnings']])
    stream=BytesIO();w.save(stream);w.close();return stream.getvalue()
