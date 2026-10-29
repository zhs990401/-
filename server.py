"""Single-user local reconciliation server. Uploaded files are never saved."""
import argparse
import base64
import csv
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO, StringIO
import json
import re
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import ZipFile, BadZipFile

import openpyxl
import xlrd
from reconcile import reconcile, reconcile_projects, reconcile_account
from cashflow import calculate_cashflow, export_cashflow
from project_import import project_template

ROOT = Path(__file__).parent / 'static'
MAX_UPLOAD = 10 * 1024 * 1024
MAX_ROWS = 30000


def cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if value is None:
        return ''
    return str(value)


def parse_file(payload):
    try:
        content = base64.b64decode(payload['content'], validate=True)
    except (ValueError, KeyError):
        raise ValueError('上传内容无效，请重新选择文件') from None
    if len(content) > MAX_UPLOAD:
        raise ValueError('每个文件不能超过 10 MB')
    suffix = Path(payload['name']).suffix.lower()
    header_row = int(payload.get('headerRow', 1))
    if not 1 <= header_row <= 50:
        raise ValueError('表头行必须在 1 至 50 行之间')
    sheets = []
    formula_count = 0
    if suffix == '.xlsx':
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Excel 解压后过大，请缩小文件')
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            sheets = workbook.sheetnames
            sheet = workbook[payload.get('sheet') or (next((s for s in sheets if '明细' in s and '余额' not in s), sheets[0]) if payload.get('autoHeader') else sheets[0])]
            rows = []
            for row in sheet.iter_rows():
                if len(rows) >= MAX_ROWS + 50:
                    raise ValueError('超过 30,000 条明细，请拆分文件后上传')
                if len(row) > 100:
                    raise ValueError('超过 100 列，请仅导出需要核对的明细')
                formula_count += sum(c.data_type == 'f' for c in row)
                rows.append([cell(c.value) for c in row])
        finally:
            workbook.close()
    elif suffix == '.xls':
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            sheets = workbook.sheet_names()
            sheet = workbook.sheet_by_name(payload.get('sheet') or (next((s for s in sheets if '明细' in s and '余额' not in s), sheets[0]) if payload.get('autoHeader') else sheets[0]))
            if sheet.nrows > MAX_ROWS + 50 or sheet.ncols > 100:
                raise ValueError('文件超过 30,000 条明细或 100 列，请拆分文件')
            rows = [[cell(xlrd.xldate_as_datetime(c.value, workbook.datemode)) if c.ctype == xlrd.XL_CELL_DATE else cell(c.value) for c in sheet.row(i)] for i in range(sheet.nrows)]
        finally:
            workbook.release_resources()
    elif suffix == '.csv':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('gb18030')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(StringIO(text), dialect):
            if len(rows) >= MAX_ROWS + 50 or len(row) > 100:
                raise ValueError('文件超过 30,000 条明细或 100 列，请拆分文件')
            rows.append(row)
    else:
        raise ValueError('仅支持 .xlsx、.xls 和 .csv 文件')
    if payload.get('autoHeader') and 'headerRow' not in payload:
        best = (0, 1)
        for i, candidate in enumerate(rows[:50]):
            names = [re.sub(r'\s+', '', str(v)) for v in candidate]
            has_date = any(n in ('日期','记账日期','业务日期','交易日期','发生日期') for n in names)
            has_amount = any(('借方' in n or '贷方' in n or n in ('原币金额','原币发生额','外币金额','外币发生额')) and '余额' not in n for n in names)
            score = int(has_date)*3 + int(has_amount)*3 + sum(any(term in n for term in ('对方项目','往来单位','币种','凭证号')) for n in names)
            if has_date and has_amount and score > best[0]:
                best = score, i+1
        header_row = best[1]
    if len(rows) < header_row:
        raise ValueError('文件中不存在该表头行')
    headers = [str(v).strip() or f'未命名列{i + 1}' for i, v in enumerate(rows[header_row - 1])]
    used_headers = set()
    for i, header in enumerate(headers):
        candidate = header
        while candidate in used_headers:
            candidate += f'（第{i + 1}列）'
        headers[i] = candidate
        used_headers.add(candidate)
    if not headers:
        raise ValueError('表头为空，请调整表头行')
    data = rows[header_row:]
    if len(data) > MAX_ROWS:
        raise ValueError('超过 30,000 条明细，请拆分文件')
    return {'headers': headers, 'rows': data, 'firstRow': header_row + 1, 'sheets': sheets,
            'sheet': (sheet.title if suffix == '.xlsx' else sheet.name if suffix == '.xls' else ''),
            'warnings': ([f'发现 {formula_count} 个公式单元格。公式不会被执行；用于核对的日期和金额列应先在 Excel 中粘贴为值。'] if formula_count else [])}


def validate_dataset(dataset):
    if not isinstance(dataset, dict) or not isinstance(dataset.get('headers'), list) or not isinstance(dataset.get('rows'), list) or not isinstance(dataset.get('mapping'), dict):
        raise ValueError('明细数据格式无效，请重新上传')
    if not 0 < len(dataset['headers']) <= 100 or len(dataset['rows']) > MAX_ROWS:
        raise ValueError('数据为空或超过处理限制')
    if not all(isinstance(row, list) and len(row) <= 100 for row in dataset['rows']):
        raise ValueError('明细行格式无效')


def export_report(report):
    """Export all categories, summaries and rules without dropping constituents."""
    if not isinstance(report, dict) or not isinstance(report.get('options'), dict):
        raise ValueError('核对报告格式无效，请先完成核对')
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    def add_sheet(name, headers, rows):
        ws = workbook.create_sheet(name)
        ws.append(headers)
        count = 0
        for row in rows:
            count += 1
            if count > 120000:
                raise ValueError('导出记录过多，请缩小核对范围')
            ws.append([str(v) if v is not None else '' for v in row])
        # Explicit string cells prevent any uploaded text becoming an Excel formula.
        for row in ws:
            for c in row:
                c.data_type = 's'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
            c.fill = openpyxl.styles.PatternFill('solid', fgColor='087A68')
        for c in ws[1]:
            ws.column_dimensions[c.column_letter].width = min(48, max(14, len(str(c.value)) * 2 + 2))

    record_headers = ['文件', '工作表', '原始行', '本方项目', '对方项目', '日期', '摘要', '参考号', '账户', '币种', '金额', '原本方名称', '原对方名称', '科目编码', '科目名称', '本币币种', '本币金额', '本币说明']

    def record_values(record):
        if not record:
            return [''] * len(record_headers)
        amount = record.get('amount')
        return [record.get(k, '') for k in ('source', 'sheet', 'row', 'project', 'counterparty', 'date', 'description', 'reference', 'account', 'currency')] + [format_cents(amount), record.get('originalProject', ''), record.get('originalCounterparty', ''), record.get('subjectCode',''), record.get('subjectName',''), record.get('baseCurrency',''), format_cents(record.get('baseAmount')), record.get('baseError','')]

    def detail_rows(entries):
        for group, entry in enumerate(entries, 1):
            left = entry.get('leftRows') or ([entry['left']] if entry.get('left') else [])
            right = entry.get('rightRows') or ([entry['right']] if entry.get('right') else [])
            for i in range(max(len(left), len(right))):
                yield [group, entry.get('pair', ''), entry.get('reason', ''), entry.get('matchType', ''), entry.get('feeStatus', '')] + record_values(left[i] if i < len(left) else None) + record_values(right[i] if i < len(right) else None) + [format_cents(entry.get('delta')) if i == 0 else '', format_cents(entry.get('baseDelta')) if i == 0 else '', entry.get('baseCurrency','')]

    for key, label in [('matched', '已匹配'), ('differences', '待核实差异'), ('ambiguous', '重复或组合冲突'), ('invalid', '无效数据'), ('fees', '疑似手续费明细'), ('translationDifferences', '本币金额差异')]:
        add_sheet(label, ['本表配对组号', '项目组合', '说明', '配对类型', '手续费核对状态'] + ['左侧' + h for h in record_headers] + ['右侧' + h for h in record_headers] + ['组金额差额', '本币金额差额', '差额本币币种'], detail_rows(report.get(key, [])))
    add_sheet('项目往来汇总', ['项目A', '项目B', '币种', 'A方净发生额', 'B方净发生额', '反向核对差额', 'A方条数', 'B方条数', '待核实项', '未提供账簿项目', '状态','金额口径','本币金额差异项数'],
              ([p['projectA'], p['projectB'], p['currency'], format_cents(p['left']), format_cents(p['right']), format_cents(p['delta']), p['leftCount'], p['rightCount'], p['outstanding'], p['missingProject'], p['status'],'原币发生额',p.get('baseDifferenceCount',0)] for p in report.get('projectPairs', [])))
    fee_keys = ['leftCount', 'leftExpense', 'leftRefund', 'rightCount', 'rightExpense', 'rightRefund', 'delta', 'missingLeftCount', 'missingLeftExpense', 'missingLeftRefund', 'missingRightCount', 'missingRightExpense', 'missingRightRefund']
    add_sheet('手续费汇总', ['币种', '左侧笔数', '左侧支出', '左侧退款', '右侧笔数', '右侧支出', '右侧退款', '规则调整后差额', '疑似左侧漏记笔数', '疑似左侧漏记支出', '疑似左侧漏记退款', '疑似右侧漏记笔数', '疑似右侧漏记支出', '疑似右侧漏记退款'],
              ([p['currency']] + [p[k] if k.endswith('Count') else format_cents(p[k]) for k in fee_keys] for p in report.get('feeSummary', [])))
    add_sheet('资金净额汇总', ['项目组合', '币种', '左侧净额', '右侧净额', '规则调整后差额'],
              ([p.get('pairLabel', ''), p['currency'], format_cents(p['left']), format_cents(p['right']), format_cents(p['delta'])] for p in report.get('totals', [])))
    add_sheet('项目导入检查', ['文件','本方项目','金额口径','所选原币金额列','币种','最早交易日','最晚交易日','记录数','无效记录数'], ([i['source'],i['project'],i['amountBasis'],' / '.join(i['amountColumns']),' / '.join(i['currency']),i['startDate'],i['endDate'],i['count'],i['invalidCount']] for i in report.get('importInventory',[])))
    add_sheet('期间外记录', ['组号','项目组合','说明','配对类型','手续费核对状态'] + ['左侧'+h for h in record_headers] + ['右侧'+h for h in record_headers] + ['原币差额','本币差额','本币币种'], detail_rows(report.get('excludedPeriod',[])))
    add_sheet('账户范围说明', ['左右侧', '文件', '工作表', '原始行', '银行账户', '未纳入原因'], ([r.get(k, '') for k in ('side','source','sheet','row','account','reason')] for r in report.get('excludedAccounts', [])))
    add_sheet('核对规则与提示', ['规则', '内容'], [[k, str(v)] for k, v in report['options'].items()] + [['项目名称对照', json.dumps(report.get('aliases', {}), ensure_ascii=False)]] + [['提示', w] for w in report.get('warnings', [])] + [['核对范围', '仅比较本次导入的发生额，不包含未导入的期初余额；疑似手续费需人工确认；一笔对多笔在唯一一对一匹配之后进行。']])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def format_cents(amount):
    if amount is None:
        return ''
    return ('-' if amount < 0 else '') + str(abs(amount) // 100) + '.' + str(abs(amount) % 100).zfill(2)


class Handler(BaseHTTPRequestHandler):
    def send(self, status, body, content_type='application/json; charset=utf-8'):
        if isinstance(body, (dict, list)):
            body = json.dumps(body, ensure_ascii=False).encode()
        elif isinstance(body, str):
            body = body.encode()
        self.send_response(status)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(body)))
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Content-Security-Policy', "default-src 'self'; script-src 'self'; style-src 'self'; connect-src 'self'; img-src 'self' data:; base-uri 'none'; frame-ancestors 'none'")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlsplit(self.path).path
        files = {'/': ('index.html', 'text/html; charset=utf-8'), '/app.js': ('app.js', 'text/javascript; charset=utf-8'), '/style.css': ('style.css', 'text/css; charset=utf-8'), '/cashflow.js': ('cashflow.js', 'text/javascript; charset=utf-8')}
        if path == '/api/project-template':
            return self.send(200, project_template(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        if path == '/api/health':
            return self.send(200, {'status': 'ok'})
        if path not in files:
            return self.send(404, {'error': '页面不存在'})
        filename, mime = files[path]
        self.send(200, (ROOT / filename).read_bytes(), mime)

    def do_POST(self):
        try:
            origin = self.headers.get('Origin')
            if origin and urlsplit(origin).netloc != self.headers.get('Host'):
                return self.send(403, {'error': '不允许跨站请求'})
            if self.headers.get('Content-Type', '').split(';')[0] != 'application/json':
                return self.send(415, {'error': '请求格式必须为 JSON'})
            length = int(self.headers.get('Content-Length', 0))
            if not 0 < length <= 40 * 1024 * 1024:
                return self.send(413, {'error': '请求为空或超过大小限制'})
            payload = json.loads(self.rfile.read(length))
            if self.path == '/api/parse':
                return self.send(200, parse_file(payload))
            if self.path == '/api/reconcile':
                for side in ('left', 'right'):
                    validate_dataset(payload[side])
                return self.send(200, reconcile_account(payload))
            if self.path == '/api/cashflow':
                datasets = payload.get('datasets')
                if not isinstance(datasets, list) or not 1 <= len(datasets) <= 20:
                    raise ValueError('现金流统计支持 1 至 20 份流水')
                for dataset in datasets:
                    validate_dataset(dataset)
                return self.send(200, calculate_cashflow(payload))
            if self.path == '/api/cashflow/export':
                return self.send(200, export_cashflow(payload), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            if self.path == '/api/projects':
                datasets = payload.get('datasets')
                if not isinstance(datasets, list) or not 1 <= len(datasets) <= 20:
                    raise ValueError('项目往来账支持 1 至 20 份文件')
                for dataset in datasets:
                    validate_dataset(dataset)
                if sum(len(d['rows']) for d in datasets) > 60000:
                    raise ValueError('本次项目往来明细合计不能超过 60,000 行，请缩小期间')
                return self.send(200, reconcile_projects(payload))
            if self.path == '/api/export':
                return self.send(200, export_report(payload), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.send(404, {'error': '接口不存在'})
        except (ValueError, KeyError, TypeError, BadZipFile, xlrd.XLRDError, openpyxl.utils.exceptions.InvalidFileException, OverflowError, csv.Error, IndexError):
            # Specific validation messages contain no file bytes or credentials.
            import sys
            error = sys.exc_info()[1]
            self.send(400, {'error': str(error)[:240] if isinstance(error, ValueError) else '无法读取文件或数据，请检查格式、工作表及字段映射'})
        except Exception:
            self.send(500, {'error': '处理失败，请确认文件未损坏且未加密，或缩小文件后重试'})


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--port', type=int, default=8000)
    args = parser.parse_args()
    server = ThreadingHTTPServer(('127.0.0.1', args.port), Handler)
    print(f'对账服务已启动，端口 {args.port}；按 Ctrl+C 停止。', flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
