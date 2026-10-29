"""Explicit original-currency import rules for project-to-project ledgers."""
from pathlib import Path
import re

from reconcile import money


def currency_name(value):
    name = re.sub(r'\s+', '', str(value or '')).upper()
    code = re.match(r'^\(?([A-Z]{3})\)?(?:[-—－_：:（(].*)?$', name)
    if code:
        return 'CNY' if code.group(1)=='RMB' else code.group(1)
    return {'人民币':'CNY','RMB':'CNY','美元':'USD','美金':'USD','US$':'USD',
            '欧元':'EUR','港币':'HKD','港元':'HKD','日元':'JPY','英镑':'GBP'}.get(name, name)


def is_base_column(name):
    return any(s in name for s in ('本币', '本位币', '折合'))


def prepare_project_dataset(dataset):
    d = dict(dataset)
    mapping = d['mapping']
    if 'projectName' not in d:
        d['projectName'] = Path(str(d.get('name', ''))).stem
    keys = ['amount'] if mapping.get('mode') == 'signed' else ['debit', 'credit']
    selected = [mapping.get(k) for k in keys if mapping.get(k)]
    if not selected:
        raise ValueError(f"{d.get('name', '项目文件')}：请选择原币金额列")
    if any(is_base_column(name) for name in selected):
        raise ValueError(f"{d.get('name', '项目文件')}：本币金额不能用于原币往来核对，请重新选择原币列")
    if any('余额' in name for name in selected):
        raise ValueError('余额列不能作为原币发生额，请导入往来辅助明细账')
    explicit = all('原币' in name or '外币' in name for name in selected)
    if not explicit and d.get('originalConfirmed') is not True:
        raise ValueError(f"{d.get('name', '项目文件')}：列名未标明原币，请先确认所选金额确为原币发生额")
    if not mapping.get('currency') and not str(d.get('currencyName', '')).strip():
        raise ValueError(f"{d.get('name', '项目文件')}：请选择原币币种列；单一币种文件也可手工填写币种，不能默认人民币")
    if mapping.get('currency') and is_base_column(mapping['currency']):
        raise ValueError('本币币种不能作为原币币种，请重新选择')
    if any(mapping.get(k) and mapping[k] not in d['headers'] for k in ('baseDebit','baseCredit','baseAmount','baseCurrency','subjectCode','subjectName')):
        raise ValueError('本币或科目辅助列不存在，请重新映射')
    d['currencyName'] = currency_name(d.get('currencyName', '')) if not mapping.get('currency') else ''
    d['amountBasis'] = 'original'
    return d


def enrich_project_rows(dataset, rows):
    mapping = dataset['mapping']
    numbers = dataset.get('rowNumbers', range(dataset.get('firstRow', 2), dataset.get('firstRow', 2) + len(dataset['rows'])))
    source = dict(zip(numbers, dataset['rows']))
    for r in rows:
        values = dict(zip(dataset['headers'], source.get(r['row'], [])))
        get = lambda key: str(values.get(mapping.get(key), '') if values.get(mapping.get(key), '') is not None else '').strip()
        r['currency'] = currency_name(r['currency'])
        r.update(subjectCode=get('subjectCode'), subjectName=get('subjectName'), baseAmount=None,
                 baseCurrency=currency_name(get('baseCurrency') if mapping.get('baseCurrency') else dataset.get('baseCurrencyName', '')),
                 baseError='', amountBasis='original')
        if any(mapping.get(k) for k in ('baseDebit','baseCredit','baseAmount')):
            try:
                if not r['baseCurrency']:
                    raise ValueError('未指定本币币种，本币金额未参与辅助比较')
                if mapping.get('baseAmount'):
                    if get('baseAmount') == '':
                        raise ValueError('本币金额为空')
                    amount = money(get('baseAmount'))
                else:
                    if get('baseDebit') == get('baseCredit') == '':
                        raise ValueError('本币借贷金额均为空')
                    debit, credit = money(get('baseDebit')), money(get('baseCredit'))
                    if debit and credit:
                        raise ValueError('本币借贷同时有金额')
                    amount = debit - credit
                r['baseAmount'] = amount
            except ValueError as error:
                r['baseError'] = str(error)
    return rows


def project_template():
    from io import BytesIO
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '往来明细'
    headers = ['本方项目','对方项目','记账日期','凭证号','摘要','科目编码','科目名称',
               '原币币种','原币借方发生额','原币贷方发生额','本币币种','本币借方发生额','本币贷方发生额']
    ws.append(headers)
    ws.freeze_panes = 'A2'
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True,color='FFFFFF')
        c.fill = openpyxl.styles.PatternFill('solid',fgColor='087A68')
        ws.column_dimensions[c.column_letter].width = 20
    notes = wb.create_sheet('填写说明')
    for row in [('项目','填写规则'),('推荐报表','往来辅助明细账，显示原币并带往来辅助核算对象'),
                ('每个项目一表','建议文件名为 项目A.xlsx；可不含本方项目列，由文件名预填后确认'),
                ('币种','每行提供明确币种，原币相同才会比较；不会默认人民币'),
                ('原币金额','填写发生额，不填余额；本币金额不能用于原币比较'),
                ('本币辅助','可保留本币金额和币种，原币一致、本币不同会另行列示，不影响原币结果'),
                ('期间','各项目导出相同期间；页面也可统一筛选起止日'),
                ('凭证号','双方账套凭证号通常不同，默认不强制相同'),
                ('期初期末余额','此版本核对发生额，不自动核对期初或期末余额'),
                ('合计行','导出逐笔记录，删除期初、本月合计、累计等汇总行')]:
        notes.append(row)
    notes.column_dimensions['A'].width=24
    notes.column_dimensions['B'].width=90
    out=BytesIO();wb.save(out);wb.close()
    return out.getvalue()
