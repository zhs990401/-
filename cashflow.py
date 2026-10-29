"""Project cash-flow movements by calendar month, with explicit opening balances."""
from collections import Counter, defaultdict
from datetime import date
import hashlib
import json
from io import BytesIO

import openpyxl
from reconcile import day, money, normalize, safe_total


def month_range(start, end):
    value = start.year * 12 + start.month - 1
    last = end.year * 12 + end.month - 1
    if last - value > 599:
        raise ValueError('单个项目统计期间不能超过 600 个月，请检查期初日期')
    while value <= last:
        year, month = divmod(value, 12)
        yield f'{year:04d}-{month + 1:02d}'
        value += 1


def calculate_cashflow(payload):
    datasets = payload.get('datasets', [])
    if not isinstance(datasets, list) or not 1 <= len(datasets) <= 20:
        raise ValueError('请导入 1 至 20 份银行流水')
    if sum(len(d['rows']) for d in datasets) > 60000:
        raise ValueError('本次流水合计不能超过 60,000 行，请按期间整理')
    openings = {}
    raw_openings = payload.get('openings', [])
    if not isinstance(raw_openings, list) or len(raw_openings) > 500:
        raise ValueError('期初设置最多 500 个项目和币种组合')
    for entry in raw_openings:
        project = str(entry.get('project', '')).strip()
        if not project:
            raise ValueError('期初设置必须填写项目名称')
        currency = str(entry.get('currency', 'CNY')).strip().upper() or 'CNY'
        key = project, currency
        if key in openings:
            raise ValueError(f'{project} / {currency} 的期初设置重复，请只保留一条')
        if str(entry.get('balance', '')).strip() == '':
            raise ValueError(f'{project} 请填写期初余额；确认为零时填写 0')
        opening = {'project': project, 'currency': currency, 'date': day(entry.get('date')).isoformat(), 'balance': money(entry['balance'])}
        openings[key] = opening
    records, invalid = [], []
    warnings = []
    for index, d in enumerate(datasets):
        if not d['mapping'].get('project'):
            raise ValueError(f"{d.get('name', '银行流水')}：请选择 project 项目列，无法从摘要推测归属")
        current = {**d, 'id': f'cashflow-{index}'}
        valid, bad = normalize(current, 'left')
        invalid.extend({'record': e['left'], 'reason': e['reason']} for e in bad)
        fingerprint = hashlib.sha256(json.dumps([d['headers'], d['rows'], d['mapping'], d.get('accountName', '')], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        for r in valid:
            r['_scope'] = r['account'] or fingerprint
        records.extend(valid)
        if not d['mapping'].get('account') and not str(d.get('accountName', '')).strip():
            warnings.append(d.get('name', '银行流水') + ' 没有账户信息，跨文件重复判断范围受限；建议填写账户或映射账户列')
    safe_total(records)
    if sum(abs(o['balance']) for o in openings.values()) + sum(abs(r['amount']) for r in records) > 9_000_000_000_000_000:
        raise ValueError('期初余额与流水总额超出精确范围，请缩小项目范围')
    options = payload.get('options', {})
    end_text = '' if options.get('endDateMode')=='unlimited' else str(options.get('endDate', '')).strip()
    dates = [day(r['date']) for r in records] + [day(o['date']) for o in openings.values()]
    if not dates:
        raise ValueError('没有可统计的日期，请检查流水和期初设置')
    end = day(end_text) if end_text else max(dates)
    for opening in openings.values():
        if day(opening['date']) > end:
            raise ValueError(f"{opening['project']} 的期初日期晚于统计截止日期")
    # Missing project is not silently assigned to another project or a zero balance.
    unclassified = [{'record': r, 'reason': 'project 为空，无法归属项目，未计入项目收支'} for r in records if not r['project']]
    eligible, excluded = [], []
    for r in records:
        if not r['project']:
            continue
        opening = openings.get((r['project'], r['currency']))
        if day(r['date']) > end:
            excluded.append({'record': r, 'reason': '晚于统计截止日期，未计入'})
        elif opening and r['date'] < opening['date']:
            excluded.append({'record': r, 'reason': '早于期初日期，未计入；期初余额应已包含这部分发生额'})
        else:
            eligible.append(r)
    # Dates + amounts + reference/account are candidates, not proof of duplication.
    def identity(r):
        common = (r['_scope'], r['currency'], r['date'], r['amount'])
        return common + (('ref', r['reference']) if r['reference'] else ('text', r['project'], r['description']))
    counts = Counter(identity(r) for r in eligible)
    include_duplicates = options.get('includeDuplicates') is True
    duplicates, accepted = [], []
    duplicate_groups = {}
    for r in eligible:
        repeated = counts[identity(r)] > 1
        if repeated:
            group = duplicate_groups.setdefault(identity(r), len(duplicate_groups) + 1)
            duplicates.append({'record': r, 'group': group, 'included': include_duplicates,
                               'reason': '疑似重复，已按用户确认全部计入' if include_duplicates else '疑似重复，暂不计入；请核实或移除重复文件后重新统计'})
        if not repeated or include_duplicates:
            accepted.append(r)
    if duplicates:
        warnings.append(f'识别 {len(duplicates)} 条疑似重复记录，' + ('用户已选择全部计入，请确认它们确为不同交易' if include_duplicates else '默认全部暂不计入；不自动认定哪一条是真实记录'))
    groups = set(openings) | {(r['project'], r['currency']) for r in records if r['project']}
    if len(groups) > 500:
        raise ValueError('项目和币种组合超过 500 个，请缩小范围')
    by_group = defaultdict(list)
    for r in accepted:
        by_group[(r['project'], r['currency'])].append(r)
    source_dates = defaultdict(list)
    for r in records:
        if r['project'] and r['date'] <= end.isoformat():
            source_dates[(r['project'], r['currency'])].append(r['date'])
    excluded_counts = Counter((e['record']['project'], e['record']['currency']) for e in excluded)
    duplicate_counts = Counter((e['record']['project'], e['record']['currency']) for e in duplicates)
    monthly, summaries = [], []
    for project, currency in sorted(groups):
        rows = by_group[(project, currency)]
        opening = openings.get((project, currency))
        relevant_dates = source_dates[(project, currency)]
        if not opening and not relevant_dates:
            continue
        start = day(opening['date']) if opening else day(min(relevant_dates))
        balance = opening['balance'] if opening else None
        movements = defaultdict(lambda: [0, 0, 0])
        for r in rows:
            current = movements[r['date'][:7]]
            current[0] += max(r['amount'], 0)
            current[1] += max(-r['amount'], 0)
            current[2] += 1
        cumulative_in, cumulative_out = 0, 0
        for month in month_range(start, end):
            income, expense, count = movements[month]
            closing = balance + income - expense if balance is not None else None
            cumulative_in += income
            cumulative_out += expense
            monthly.append({'project': project, 'currency': currency, 'month': month, 'opening': balance,
                            'income': income, 'expense': expense, 'net': income - expense, 'closing': closing,
                            'cumulativeIncome': cumulative_in, 'cumulativeExpense': cumulative_out, 'count': count,
                            'periodStart': start.isoformat(), 'periodEnd': end.isoformat()})
            balance = closing
        excluded_count = excluded_counts[(project, currency)]
        duplicate_count = duplicate_counts[(project, currency)]
        summaries.append({'project': project, 'currency': currency, 'startDate': start.isoformat(), 'endDate': end.isoformat(),
                          'opening': opening['balance'] if opening else None, 'income': cumulative_in, 'expense': cumulative_out,
                          'net': cumulative_in - cumulative_out, 'closing': balance, 'count': len(rows),
                          'excludedCount': excluded_count, 'duplicateCount': duplicate_count,
                          'status': '缺少期初余额，仅统计收支' if not opening else '存在疑似重复待核实' if duplicate_count and not include_duplicates else '按已导入数据统计'})
        if not opening:
            warnings.append(f'{project} / {currency} 缺少期初设置，只统计收支，不计算余额')
        if len(monthly) > 120000:
            raise ValueError('项目月份行数超过 120,000，请缩小期间')
    if invalid or unclassified:
        warnings.append(f'无效记录 {len(invalid)} 条、未分类记录 {len(unclassified)} 条未纳入统计，请核实')
    warnings.append('按已导入银行流水统计，包含账户间调拨；未自动抵销内部转账。无流水的月份仅表示没有导入记录，不证明收支为零或数据齐全。')
    # Internal duplicate scope is never part of exported financial data.
    for r in records:
        r.pop('_scope', None)
    totals = []
    for currency in sorted({s['currency'] for s in summaries}):
        current = [s for s in summaries if s['currency'] == currency]
        totals.append({'currency': currency, 'income': sum(s['income'] for s in current), 'expense': sum(s['expense'] for s in current),
                       'net': sum(s['net'] for s in current), 'closing': sum(s['closing'] for s in current) if all(s['closing'] is not None for s in current) else None})
    return {'mode': 'cashflow', 'summaries': summaries, 'monthly': monthly, 'transactions': accepted,
            'openings': list(openings.values()), 'invalid': invalid, 'unclassified': unclassified, 'duplicates': duplicates,
            'excluded': excluded, 'totals': totals, 'warnings': list(dict.fromkeys(warnings)),
            'options': {'endDateMode':options.get('endDateMode','custom' if end_text else 'unlimited'),'endDate': end.isoformat(), 'includeDuplicates': include_duplicates},
            'counts': {'files': len(datasets), 'input': len(records) + len(invalid), 'included': len(accepted), 'projects': len({s['project'] for s in summaries})}}


def export_cashflow(report):
    if report.get('mode') != 'cashflow':
        raise ValueError('现金流报告格式无效')
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    def amount(value):
        if value is None:
            return ''
        return ('-' if value < 0 else '') + f'{abs(value)//100}.{abs(value)%100:02d}'
    def sheet(name, columns, rows, money_keys=()):
        ws = wb.create_sheet(name)
        ws.append([label for _, label in columns])
        for row in rows:
            ws.append([amount(row.get(key)) if key in money_keys else str(row.get(key, '')) for key, _ in columns])
        for row in ws:
            for c in row:
                c.data_type = 's'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
            c.fill = openpyxl.styles.PatternFill('solid', fgColor='087A68')
            ws.column_dimensions[c.column_letter].width = 20
    money_keys = ('opening', 'income', 'expense', 'net', 'closing', 'cumulativeIncome', 'cumulativeExpense', 'balance', 'amount')
    base = [('project','项目'),('currency','币种')]
    sheet('项目全周期汇总', base + [('startDate','统计起始日'),('endDate','统计截止日'),('opening','期初余额'),('income','累计收入'),('expense','累计支出'),('net','累计净额'),('closing','期末余额'),('count','计入笔数'),('status','统计状态')], report['summaries'], money_keys)
    sheet('项目月度现金流', base + [('month','月份'),('opening','月初余额'),('income','当月收入'),('expense','当月支出'),('net','当月净额'),('closing','月末余额'),('cumulativeIncome','累计收入'),('cumulativeExpense','累计支出'),('count','计入笔数'),('periodStart','期初日期'),('periodEnd','统计截止日')], report['monthly'], money_keys)
    sheet('手工期初设置', base + [('date','期初日期'),('balance','期初余额')], report['openings'], money_keys)
    detail = base + [('source','文件'),('sheet','工作表'),('row','原始行'),('account','银行账户'),('date','日期'),('description','摘要'),('reference','流水号'),('amount','金额'),('reason','说明'),('group','重复候选组'),('included','用户确认计入')]
    sheet('已计入流水', detail, report['transactions'], money_keys)
    for key, name in [('invalid','无效记录'),('unclassified','未分类记录'),('duplicates','疑似重复记录'),('excluded','期间范围外记录')]:
        sheet(name, detail, [{**e['record'], **{k:v for k,v in e.items() if k != 'record'}} for e in report[key]], money_keys)
    sheet('统计规则与提示', [('key','规则'),('value','内容')], [{'key':k,'value':str(v)} for k,v in report['options'].items()] + [{'key':'提示','value':w} for w in report['warnings']])
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()
