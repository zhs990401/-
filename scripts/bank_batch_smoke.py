"""Public-browser multi-file/account regression using fictional workbooks only."""
import base64
from io import BytesIO
import importlib.util
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
from urllib.request import urlopen
from urllib.error import URLError

import openpyxl
from playwright.sync_api import sync_playwright, expect

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
spec=importlib.util.spec_from_file_location('bank_fixture',ROOT/'tests/test_bank_batch.py')
fixture=importlib.util.module_from_spec(spec);spec.loader.exec_module(fixture)
def upload(payload):
    return dict(name=payload['name'],mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',buffer=base64.b64decode(payload['content']))
with socket.socket() as s:s.bind(('127.0.0.1',0));port=s.getsockname()[1]
settings={**os.environ,'PORT':str(port)}
process=subprocess.Popen([str(ROOT/'.venv/bin/gunicorn'),'--config','gunicorn.conf.py','webapp:create_app()'],cwd=ROOT,env=settings,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
try:
    base=f'http://127.0.0.1:{port}'
    for _ in range(60):
        if process.poll() is not None:raise RuntimeError('Gunicorn failed')
        try:
            with urlopen(base+'/api/health',timeout=1):break
        except (URLError,TimeoutError):time.sleep(.2)
    else:raise RuntimeError('Gunicorn not ready')
    with sync_playwright() as p,tempfile.TemporaryDirectory() as output:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
        page=browser.new_page(viewport={'width':1440,'height':1000});errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        page.goto(base+'/#bankbatch')
        expect(page).to_have_url(base+'/#bank')
        expect(page.locator('#nav-bankbatch')).to_have_count(0)
        expect(page.locator('#nav-bank')).to_have_attribute('aria-current','page')
        expect(page.locator('#bankbatch-workspace')).to_be_visible()
        page.locator('#bb-left').set_input_files(upload(fixture.inner_file()))
        expect(page.locator('#bb-import-check')).to_contain_text('12345678')
        # Upload two external files together, each containing its own bank Sheet.
        outer=fixture.outer_file();split=[]
        w=openpyxl.load_workbook(BytesIO(base64.b64decode(outer['content'])))
        titles=list(w.sheetnames);w.close()
        for i,title in enumerate(titles):
            w=openpyxl.load_workbook(BytesIO(base64.b64decode(outer['content'])))
            for other in list(w.sheetnames):
                if other!=title:del w[other]
            split.append(upload(fixture.encode(w,f'Account Summary 01 Jan 2026 - 31 Jan 2026 TEST{i}.xlsx')))
        page.locator('#bb-right').set_input_files(split)
        expect(page.locator('#bb-right-files p')).to_have_count(2)
        expect(page.locator('#bb-start')).to_have_value('2026-01-01')
        expect(page.locator('#bb-end')).to_have_value('2026-01-31')
        page.locator('#bb-run').click()
        expect(page.locator('#bb-summary')).to_contain_text('期末差异由手续费解释')
        expect(page.locator('#bb-summary')).to_contain_text('缺少内账')
        expect(page.locator('#bb-fees-summary')).to_contain_text('136.00')
        page.locator('#bb-account').select_option('0')
        expect(page.locator('#bb-summary tr')).to_have_count(1)
        page.get_by_role('button',name='手续费（含透支利息） 4',exact=True).click()
        expect(page.locator('#bb-details')).to_contain_text('透支利息')
        expect(page.locator('#bb-details tr')).to_have_count(2)
        with page.expect_download() as info:page.locator('#bb-export').click()
        dest=Path(output)/'report.xlsx';info.value.save_as(dest)
        w=openpyxl.load_workbook(dest)
        assert w['账户核对总览'].max_row==3
        assert w['外账余额调节项目'].max_row==2
        w.close()
        # Append the second internal account without losing the first account.
        payload=fixture.inner_file();w=openpyxl.load_workbook(BytesIO(base64.b64decode(payload['content'])))
        for row in w.active.iter_rows(min_row=5):
            row[5].value='87654321';row[4].value='USD'
        page.locator('#bb-left').set_input_files(upload(fixture.encode(w,'第二账户.xlsx')))
        expect(page.locator('#bb-left-files p')).to_have_count(2)
        page.locator('#bb-run').click()
        expect(page.locator('#bb-summary tr')).to_have_count(2)
        expect(page.locator('#bb-summary')).not_to_contain_text('缺少内账')
        page.get_by_role('button',name='手续费（含透支利息） 4',exact=True).click()
        page.locator('[data-bb-classify]').first.click()
        expect(page.locator('#bb-results')).to_be_hidden()
        page.locator('#bb-run').click()
        expect(page.locator('#bb-tabs')).to_contain_text('手续费（含透支利息） 3')
        page.locator('#nav-projects').click();page.locator('#nav-bank').click()
        expect(page.locator('#bb-summary tr')).to_have_count(2)
        page.locator('#bank-format-options summary').click()
        page.locator('#bank-manual-open').click()
        expect(page.locator('#reconciliation-workspace')).to_be_visible()
        expect(page.locator('#bankbatch-workspace')).to_be_hidden()
        page.locator('#bank-batch-open').click()
        expect(page.locator('#bb-summary tr')).to_have_count(2)
        expect(page).to_have_url(base+'/#bank')
        page.set_viewport_size({'width':390,'height':844})
        expect(page.locator('#bb-account')).to_be_visible()
        assert not errors,errors
        browser.close()
        print('Batch bank browser passed: public entry, two external files/Sheets, multiple internal files, account/currency filters, balances, overdraft fees/refunds, missing counterpart, adjustment isolation, Excel export, manual classification, module persistence and mobile layout.')
finally:
    process.terminate()
    try:process.wait(timeout=10)
    except subprocess.TimeoutExpired:process.kill();process.wait(timeout=5)
