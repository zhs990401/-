"""Optional real-browser checks. Run server.py first; requires Playwright + Chromium."""
from io import BytesIO
from pathlib import Path
import tempfile
from datetime import datetime
import csv
import os
import openpyxl
from playwright.sync_api import sync_playwright, expect


def xlsx(headers, rows):
    wb = openpyxl.Workbook()
    wb.active.title = '往来明细'
    wb.active.append(headers)
    for row in rows:
        wb.active.append(row)
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def upload(name, content):
    return {'name': name, 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': content}


with sync_playwright() as p, tempfile.TemporaryDirectory(prefix='reconcile-browser-') as output:
    browser = p.chromium.launch(executable_path='/usr/bin/chromium', headless=True, args=['--no-sandbox'])
    credentials = None
    if os.environ.get('TEST_APP_USERNAME') and os.environ.get('TEST_APP_PASSWORD'):
        credentials = {'username': os.environ['TEST_APP_USERNAME'], 'password': os.environ['TEST_APP_PASSWORD']}
    page = browser.new_page(viewport={'width':1440,'height':1100}, device_scale_factor=1, http_credentials=credentials)
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.goto(os.environ.get('TEST_BASE_URL', 'http://127.0.0.1:8000'))
    expect(page.locator('.brand-name')).to_have_text('中国港湾账目核对平台')
    expect(page.locator('#nav-bank')).to_have_attribute('aria-current', 'page')
    expect(page.locator('#feature-title')).to_have_text('银行资金内外账核对')
    expect(page.locator('#project-demo')).to_be_hidden()
    page.locator('#demo').click()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents() == ['3','3','3','1','2']
    expect(page.locator('#fee-summary')).to_contain_text('12.00 / 2.00')
    page.get_by_role('button', name='已匹配 3', exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('一笔对多笔')
    expect(page.locator('#report-body')).to_contain_text('材料付款第二笔')
    page.get_by_role('button', name='疑似手续费 2', exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('疑似内账漏记')
    with page.expect_download() as info:
        page.locator('#export-excel').click()
    dest = Path(output)/'bank.xlsx'
    info.value.save_as(dest)
    wb = openpyxl.load_workbook(dest)
    assert wb['已匹配'].max_row == 5  # 3 groups, including a group with 2 components.
    assert wb['疑似手续费明细'].max_row == 3
    assert wb['手续费汇总'].max_row == 2
    wb.close()
    with page.expect_download() as info:
        page.locator('#export').click()
    dest = Path(output)/'details.csv'
    info.value.save_as(dest)
    assert '材料付款第二笔' in dest.read_text(encoding='utf-8-sig')
    # Each sidebar module retains its own rules, search, data and result.
    page.locator('#search').fill('手续费')
    page.locator('#nav-projects').click()
    expect(page.locator('#feature-title')).to_have_text('项目往来核对')
    expect(page.locator('#bank-workflow')).to_be_hidden()
    expect(page.locator('#fee-rule')).to_be_hidden()
    expect(page.locator('#results')).to_be_hidden()
    page.locator('#project-demo').click()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents() == ['2','2','0','0','1'], page.locator('.stat strong').all_text_contents()
    expect(page.locator('#project-summary tbody tr')).to_have_count(4)
    expect(page.locator('#report-warnings')).to_contain_text('项目D')
    expect(page.locator('#direction')).to_be_disabled()
    with page.expect_download() as info:
        page.locator('#export-excel').click()
    dest = Path(output)/'projects.xlsx'
    info.value.save_as(dest)
    wb = openpyxl.load_workbook(dest)
    assert wb['项目往来汇总'].max_row == 5
    assert wb['本币金额差异'].max_row == 3  # One group, two right-side constituents.
    wb.close()
    page.locator('#project-pair-filter').select_option(label='项目A ↔ 项目B')
    expect(page.locator('#project-summary tbody tr')).to_have_count(1)
    page.get_by_role('button',name='本币金额差异 1',exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('CNY 本币差额 100.00')
    page.locator('#project-pair-filter').select_option('')
    with page.expect_download() as info:
        page.get_by_role('link',name='下载导表模板').click()
    dest=Path(output)/'template.xlsx';info.value.save_as(dest)
    wb=openpyxl.load_workbook(dest)
    assert '原币借方发生额' in [c.value for c in wb['往来明细'][1]]
    wb.close()
    page.locator('#nav-bank').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents() == ['3','3','3','1','2']
    expect(page.locator('#search')).to_have_value('手续费')
    page.locator('#days').select_option('3')
    expect(page.locator('#results')).to_be_hidden()
    page.locator('#nav-projects').click()
    expect(page.locator('#days')).to_have_value('0')
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents() == ['2','2','0','0','1'], page.locator('.stat strong').all_text_contents()
    page.locator('#nav-bank').click()
    expect(page.locator('#days')).to_have_value('3')
    expect(page.locator('#results')).to_be_hidden()
    page.locator('#nav-projects').click()
    page.evaluate('window.scrollTo(0,0)')
    page.screenshot(path='/tmp/衡账-多项目核对.png', full_page=True)
    page.set_viewport_size({'width':390,'height':844})
    assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
    page.screenshot(path='/tmp/衡账-手机.png', full_page=True)
    page.set_viewport_size({'width':1440,'height':1100})

    # Actual multi-file XLSX upload, explicit project names, and missing project.
    page.reload()
    expect(page.locator('#nav-projects')).to_have_attribute('aria-current', 'page')
    page.locator('#nav-projects').click()
    headers = ['日期','对方项目','摘要','原币借方发生额','原币贷方发生额','参考号','原币币种']
    a = xlsx(headers, [[datetime(2026,1,1),'项目B','拨付',1000,None,'AB','USD']])
    b = xlsx(headers, [[datetime(2026,1,1),'项目A','第一笔',None,600,'AB','USD'],[datetime(2026,1,1),'项目A','第二笔',None,400,'AB','USD']])
    c = xlsx(headers, [[datetime(2026,1,1),'项目D','缺少对方',20,None,'CD','USD']])
    page.locator('#project-batch').set_input_files([upload('项目A.xlsx',a),upload('项目B.xlsx',b),upload('项目C.xlsx',c)])
    expect(page.locator('#file-p3').locator('..').locator('.dropzone b')).to_have_text('项目C.xlsx')
    for id_, name in [('p1','项目A'),('p2','项目B'),('p3','项目C')]:
        expect(page.locator('[data-project-name="'+id_+'"]')).to_have_value(name)
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','1','0','0','0']
    expect(page.locator('#report-warnings')).to_contain_text('项目D')

    # Combined sheet uses project-column mapping; aliases resolve both directions.
    page.reload()
    page.locator('#nav-projects').click()
    combined=xlsx(['日期','本方项目','对方项目','借方金额','贷方金额','摘要'],[
        [datetime(2026,1,1),'甲','乙',100,None,'<script>保留为文本</script>'],
        [datetime(2026,1,1),'B','A',None,100,'反向']])
    page.locator('#file-p1').set_input_files(upload('合并表.xlsx',combined))
    expect(page.locator('#file-p1').locator('..').locator('.dropzone b')).to_have_text('合并表.xlsx')
    page.locator('#aliases').fill('甲=A\n乙=B')
    expect(page.locator('#run')).to_be_disabled()
    page.locator('[data-project-setting="originalConfirmed"]').check()
    page.locator('[data-project-setting="currencyName"]').fill('CNY')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','0','0','0','0']
    page.get_by_role('button',name='已匹配 1',exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('<script>保留为文本</script>')
    expect(page.locator('#report-body script')).to_have_count(0)
    page.locator('#days').select_option('1')
    expect(page.locator('#results')).to_be_hidden()
    # Real exports with a title row, base columns before original columns, and different vouchers.
    page.reload()
    original_headers=['记账日期','往来单位名称','本币借方发生额','本币贷方发生额','原币借方发生额','原币贷方发生额','原币币种','本币币种','凭证号']
    def titled_export(rows):
        wb=openpyxl.Workbook();wb.active.title='说明';wb.active.append(['导出说明'])
        sheet=wb.create_sheet('往来明细');sheet.append(['内部往来辅助明细账']);sheet.append(original_headers)
        for row in rows:sheet.append(row)
        stream=BytesIO();wb.save(stream);wb.close();return stream.getvalue()
    a=titled_export([[datetime(2025,12,31),'项目B',72,None,10,None,'USD','CNY','A0'],[datetime(2026,1,1),'项目B',7200,None,1000,None,'USD','CNY','A1']])
    b=titled_export([[datetime(2026,1,1),'项目A',None,7300,None,1000,'美元','CNY','B99']])
    page.locator('#project-batch').set_input_files([upload('项目A.xlsx',a),upload('项目B.xlsx',b)])
    expect(page.locator('#file-p2').locator('..').locator('.dropzone b')).to_have_text('项目B.xlsx')
    expect(page.locator('[data-id="p1"][data-map="debit"]')).to_have_value('原币借方发生额')
    expect(page.locator('#header-p1')).to_have_value('2')
    expect(page.locator('#sheet-p1')).to_have_value('往来明细')
    expect(page.locator('#project-import-check')).to_contain_text('0 份待确认')
    page.locator('#project-date-range').select_option('custom')
    page.locator('#project-start').fill('2026-01-01')
    page.locator('#project-end').fill('2026-01-31')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','0','0','0','1']
    expect(page.locator('#report-warnings')).to_contain_text('1 条记录不在选定期间')
    page.get_by_role('button',name='本币金额差异 1',exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('-100.00')
    page.locator('[data-id="p1"][data-map="debit"]').select_option('本币借方发生额')
    expect(page.locator('#run')).to_be_disabled()
    expect(page.locator('#project-import-check')).to_contain_text('所选金额不是原币发生额')

    # An explicitly confirmed original-currency mapping can be shared by identical templates.
    page.reload()
    generic_headers=['日期','对方项目','借方金额','贷方金额']
    a=xlsx(generic_headers,[[datetime(2026,1,1),'项目B',10,None]])
    b=xlsx(generic_headers,[[datetime(2026,1,1),'项目A',None,10]])
    page.locator('#project-batch').set_input_files([upload('项目A.xlsx',a),upload('项目B.xlsx',b)])
    expect(page.locator('#file-p2').locator('..').locator('.dropzone b')).to_have_text('项目B.xlsx')
    page.locator('[data-id="p1"][data-project-setting="originalConfirmed"]').check()
    page.locator('[data-id="p1"][data-project-setting="currencyName"]').fill('USD')
    expect(page.locator('#run')).to_be_disabled()
    page.locator('[data-project-apply="p1"]').click()
    expect(page.locator('[data-project-name="p2"]')).to_have_value('项目B')
    expect(page.locator('[data-id="p2"][data-project-setting="currencyName"]')).to_have_value('USD')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','0','0','0','0']

    # One combined 单位三栏账 discovers books and splits all internal units automatically.
    page.reload()
    import importlib.util
    spec=importlib.util.spec_from_file_location('fictional_unit_fixture',Path(__file__).resolve().parents[1]/'tests/unit_ledger_fixture.py')
    fixture=importlib.util.module_from_spec(spec);spec.loader.exec_module(fixture)
    content=fixture.fixture_bytes()
    page.locator('#project-consolidated').set_input_files(upload('单位三栏账.xlsx',content))
    expect(page.locator('#unit-ledger-discovery')).to_contain_text('3 个内部单位')
    expect(page.locator('#uploads .upload-card')).to_have_count(1)
    expect(page.locator('#days')).to_have_value('unlimited')
    expect(page.locator('#project-date-range')).to_have_value('unlimited')
    expect(page.locator('#project-start')).to_be_disabled()
    expect(page.locator('#run')).to_be_enabled()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    expect(page.locator('#unit-project-overview')).to_contain_text('甲项目')
    expect(page.locator('#unit-project-overview')).to_contain_text('乙项目')
    expect(page.locator('#unit-project-overview')).to_contain_text('丙项目')
    expect(page.locator('#unit-project-overview tbody tr')).to_have_count(4)
    expect(page.locator('#self-project-summary')).to_be_visible()
    expect(page.locator('#self-project-summary')).to_contain_text('本项目自挂往来核对')
    expect(page.locator('#self-project-summary')).to_contain_text('40.00')
    page.get_by_role('button',name='未纳入记录 5',exact=True).click()
    expect(page.locator('#report-body tr')).to_have_count(5)
    page.get_by_role('button',name='已匹配 1',exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('日期无限制')
    page.locator('[data-open-unit="甲项目"]').first.click()
    expect(page.locator('#project-unit-filter')).to_have_value('甲项目')
    expect(page.locator('#unit-project-overview tbody tr')).to_have_count(2)
    with page.expect_download() as info:
        page.locator('#export-excel').click()
    dest=Path(output)/'combined.xlsx';info.value.save_as(dest)
    wb=openpyxl.load_workbook(dest)
    assert wb['内部单位清单'].max_row==4
    assert wb['未纳入内部核对'].max_row==6
    wb.close()
    page.locator('#project-unit-filter').select_option('')
    page.locator('#project-consolidated').set_input_files(upload('单位三栏账.xlsx',content))
    expect(page.locator('#notice')).to_contain_text('当前已有项目文件')
    expect(page.locator('#uploads .upload-card')).to_have_count(1)
    expect(page.locator('#results')).to_be_visible()
    page.locator('#days').select_option('0')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').first.inner_text()=='0'
    page.locator('#days').select_option('unlimited')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').first.inner_text()=='1'
    page.evaluate('window.scrollTo(0,0)')
    page.screenshot(path='/tmp/中国港湾-总表自动核对-虚构数据.png',full_page=True)

    # Bank reconciliation is constrained to one explicit account, including report totals.
    page.locator('#nav-bank').click()
    bank_headers=['日期','账户','收入金额','支出金额','流水号']
    one=xlsx(bank_headers,[[datetime(2026,1,1),'BANK-A',100,None,'A1'],[datetime(2026,1,2),'BANK-B',20,None,'B1']])
    two=xlsx(bank_headers,[[datetime(2026,1,1),'BANK-A',90,None,'A1'],[datetime(2026,1,2),'BANK-B',20,None,'B1']])
    for side, content in [('left',one),('right',two)]:
        page.locator('#file-'+side).set_input_files(upload(side+'.xlsx',content))
        expect(page.locator('#file-'+side).locator('..').locator('.dropzone b')).to_have_text(side+'.xlsx')
    page.locator('#accountScope').select_option('BANK-B')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','0','0','0','0']
    expect(page.locator('#totals')).to_contain_text('20.00')
    expect(page.locator('#result-summary')).to_contain_text('BANK-B')
    page.locator('#accountScope').select_option('BANK-A')
    expect(page.locator('#results')).to_be_hidden()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['0','1','0','0','0']

    page.locator('#days').select_option('unlimited')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()

    # Repeated equal amounts match by count without inventing transaction pairings.
    count_headers=['日期','账户','收入金额','支出金额','流水号']
    left_repeats=xlsx(count_headers,[[datetime(2026,1,1),'COUNT-A',100,None,''],[datetime(2026,1,1),'COUNT-A',100,None,'']])
    right_repeats=xlsx(count_headers,[[datetime(2026,3,1),'COUNT-A',100,None,''],[datetime(2026,3,1),'COUNT-A',100,None,'']])
    for side,content in [('left',left_repeats),('right',right_repeats)]:
        page.locator('#file-'+side).set_input_files(upload('count_'+side+'.xlsx',content))
        expect(page.locator('#file-'+side).locator('..').locator('.dropzone b')).to_have_text('count_'+side+'.xlsx')
    page.locator('#accountScope').select_option('COUNT-A')
    page.locator('#days').select_option('unlimited')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','0','0','0','0']
    expect(page.locator('#duplicate-count-summary')).to_contain_text('金额及数量一致，按组匹配')
    expect(page.locator('#result-summary')).to_contain_text('1 组金额数量匹配')
    page.get_by_role('button',name='已匹配 1',exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('组内未指定逐笔对应')
    expect(page.locator('#report-body')).to_contain_text('2 笔')
    with page.expect_download() as info:
        page.locator('#export-excel').click()
    dest=Path(output)/'quantity.xlsx';info.value.save_as(dest)
    wb=openpyxl.load_workbook(dest)
    assert wb['已匹配'].max_row==5
    assert wb['重复金额数量核对'].max_row==2
    columns=[c.value for c in wb['已匹配'][1]]
    li,ri=columns.index('左侧文件'),columns.index('右侧文件')
    assert all(bool(r[li])!=bool(r[ri]) for r in wb['已匹配'].iter_rows(min_row=2,values_only=True))
    wb.close()
    page.locator('#duplicate-counts').uncheck()
    expect(page.locator('#results')).to_be_hidden()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['0','0','4','0','0']
    page.locator('#duplicate-counts').check()
    right_one=xlsx(count_headers,[[datetime(2026,3,1),'COUNT-A',100,None,'']])
    page.locator('#file-right').set_input_files(upload('count_right_one.xlsx',right_one))
    expect(page.locator('#file-right').locator('..').locator('.dropzone b')).to_have_text('count_right_one.xlsx')
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['0','0','3','0','0']
    expect(page.locator('#duplicate-count-summary')).to_contain_text('数量不一致，人工复核')
    expect(page.locator('#duplicate-count-summary')).to_contain_text('原始左 2 笔 / 右 1 笔')

    # Repeated split amounts form one group, with separate source rows in exports.
    repeated_parts=xlsx(count_headers,[[datetime(2026,3,1),'COUNT-A',amount,None,''] for amount in [60,40,60,40]])
    page.locator('#file-right').set_input_files(upload('repeated_parts.xlsx',repeated_parts))
    expect(page.locator('#file-right').locator('..').locator('.dropzone b')).to_have_text('repeated_parts.xlsx')
    page.locator('#oneToMany').check()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['1','0','0','0','0']
    expect(page.locator('#result-summary')).to_contain_text('1 组重复一对多按组匹配')
    expect(page.locator('#duplicate-count-summary')).to_be_hidden()
    page.get_by_role('button',name='已匹配 1',exact=True).click()
    expect(page.locator('#report-body')).to_contain_text('左 2 笔 ↔ 右 4 笔')
    expect(page.locator('#report-body')).to_contain_text('组内未指定逐笔对应')
    expect(page.locator('#report-body')).to_contain_text('40.00 × 2；60.00 × 2')
    expect(page.locator('#report-body')).to_contain_text('合计 200.00')
    for details in page.locator('#report-body details').all():
        details.locator('summary').click()
    expect(page.locator('#report-body')).to_contain_text('第 5 行')
    with page.expect_download() as info:
        page.locator('#export-excel').click()
    dest=Path(output)/'repeated_split.xlsx';info.value.save_as(dest)
    wb=openpyxl.load_workbook(dest)
    assert wb['已匹配'].max_row==7
    assert wb['重复或组合冲突'].max_row==1
    columns=[c.value for c in wb['已匹配'][1]]
    li,ri=columns.index('左侧文件'),columns.index('右侧文件')
    assert all(bool(r[li])!=bool(r[ri]) for r in wb['已匹配'].iter_rows(min_row=2,values_only=True))
    wb.close()
    with page.expect_download() as info:
        page.locator('#export').click()
    dest=Path(output)/'repeated_split.csv';info.value.save_as(dest)
    with dest.open(encoding='utf-8-sig',newline='') as stream:
        exported=list(csv.DictReader(stream))
    members=[r for r in exported if r['分类']=='已匹配']
    assert len(members)==6
    assert all(bool(r['左侧文件'])!=bool(r['右侧文件']) for r in members)
    assert len({r['组号'] for r in members})==1
    page.locator('#oneToMany').uncheck()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()
    assert page.locator('.stat strong').all_text_contents()==['0','0','6','0','0']
    page.locator('#oneToMany').check()
    page.locator('#run').click()
    expect(page.locator('#results')).to_be_visible()

    # Cash-flow demo, project/month totals, graph, reports, and per-module persistence.
    page.locator('#nav-cashflow').click()
    expect(page.locator('#feature-title')).to_have_text('项目现金流统计')
    expect(page.locator('#reconciliation-workspace')).to_be_hidden()
    page.locator('#cashflow-demo').click()
    page.locator('#cf-run').click()
    expect(page.locator('#cf-results')).to_be_visible()
    expect(page.locator('#cf-summary-body')).to_contain_text('1,100.00')
    expect(page.locator('#cf-summary-body')).to_contain_text('700.00')
    expect(page.locator('#cf-warnings')).to_contain_text('未分类记录 1 条')
    expect(page.locator('#cf-monthly-body tr')).to_have_count(6)
    page.locator('#cf-end-policy').select_option('unlimited')
    expect(page.locator('#cf-end')).to_be_disabled()
    page.locator('#cf-run').click()
    expect(page.locator('#cf-results')).to_be_visible()
    expect(page.locator('#cf-result-description')).to_contain_text('2026-03-03')
    page.locator('#cf-end-policy').select_option('custom')
    page.locator('#cf-end').fill('2026-03-31')
    page.locator('#cf-run').click()
    expect(page.locator('#cf-results')).to_be_visible()
    page.locator('#cf-project-filter').select_option(label='项目A / CNY')
    expect(page.locator('#cf-chart svg')).to_be_visible()
    expect(page.locator('#cf-monthly-body tr')).to_have_count(3)
    with page.expect_download() as info:
        page.locator('#cf-export').click()
    dest=Path(output)/'cashflow.xlsx';info.value.save_as(dest)
    wb=openpyxl.load_workbook(dest)
    assert wb['项目月度现金流'].max_row==7
    assert wb['未分类记录'].max_row==2
    wb.close()
    with page.expect_download() as info:
        page.locator('#cf-save').click()
    work=Path(output)/'work.json';info.value.save_as(work)
    page.locator('#nav-bank').click()
    expect(page.locator('#accountScope')).to_have_value('COUNT-A')
    expect(page.locator('#results')).to_be_visible()
    page.locator('#nav-cashflow').click()
    expect(page.locator('#cf-results')).to_be_visible()
    expect(page.locator('#cf-project-filter')).to_have_value('["项目A","CNY"]')
    page.evaluate('window.scrollTo(0,0)')
    page.screenshot(path='/tmp/中国港湾-现金流统计.png',full_page=True)
    page.set_viewport_size({'width':390,'height':844})
    assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
    page.set_viewport_size({'width':1440,'height':1100})

    # Restore a local workpack after a page reload, then append a real XLSX month.
    page.reload()
    expect(page.locator('#nav-cashflow')).to_have_attribute('aria-current','page')
    page.locator('#cf-restore').set_input_files(str(work))
    expect(page.locator('#cf-notice')).to_contain_text('工作包已读取')
    page.locator('#cf-run').click()
    expect(page.locator('#cf-results')).to_be_visible()
    expect(page.locator('#cf-summary-body')).to_contain_text('1,100.00')
    april=xlsx(['日期','project','收入金额','支出金额','流水号','账户'],[[datetime(2026,4,2),'项目A',250,None,'T8','银行001']])
    page.locator('#cf-files').set_input_files(upload('四月流水.xlsx',april))
    expect(page.locator('#cf-uploads')).to_contain_text('四月流水.xlsx')
    page.locator('#cf-end').fill('2026-04-30')
    page.locator('#cf-run').click()
    expect(page.locator('#cf-results')).to_be_visible()
    expect(page.locator('#cf-summary-body')).to_contain_text('1,350.00')
    # Uploading the same month twice becomes a visible exception, never silent double count.
    page.locator('#cf-files').set_input_files(upload('四月流水再上传.xlsx',april))
    expect(page.locator('#cf-uploads')).to_contain_text('四月流水再上传.xlsx')
    page.locator('#cf-run').click()
    expect(page.locator('#cf-results')).to_be_visible()
    expect(page.locator('#cf-warnings')).to_contain_text('2 条疑似重复')
    expect(page.locator('#cf-summary-body')).to_contain_text('1,100.00')
    page.get_by_role('button',name='疑似重复 2',exact=True).click()
    expect(page.locator('#cf-details tr')).to_have_count(2)
    assert not errors, errors
    browser.close()
    print('Browser checks passed: China Harbour platform title, separate sidebar modules with independent data/rules/results/search, URL restoration, bank fee/refund stats, group members, XLSX/CSV export, batch upload, aliases, missing project warnings, HTML escaping, mobile layout, single-bank-account scope, project/month cashflow, graphs, local workpack restore and duplicate upload review; original/base currency separation, combined three-column ledger upload, automatic internal-unit discovery, project overview, explicit unlimited date controls, auto headers, mapping reuse, project/currency filters and template download; repeated-amount equal/unequal counts, repeated split patterns, switches, constituent counts and non-pairwise XLSX/CSV group export; no JavaScript errors.')
